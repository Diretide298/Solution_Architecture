#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Every question between the package and 100% ready on every module, in one workbook.

Asked for on 28 September. The readiness report says *what* stands between each module and 100%;
this is the list somebody takes into the room. Nothing here is authored. It is read from the
contracts, the screens, the design-pack index and the audit decisions, so it shrinks as answers land.

Sheets:
    Summary       each platform: ready now, and which questions close the gap
    Sessions      the five client sessions and the scoping decision, with their size
    Flagged       operations whose drafted shape needs a real answer, not a yes (one question each)
    Sign-off      the rest of the unagreed operations: agree as drafted, correct, or not needed
    Scope         design packs nothing has been drafted from: in this delivery or not?
    Specify       screens that name no operation: what does each show and call?
    Wireframes    each platform's frames: who verifies, which revision
    Open values   values the audit decisions left for the client to name

    python3 tools/build-readiness-questions.py
Writes handoff/TICVAI_Readiness_Questions.xlsx.
"""
import collections
import glob
import importlib.util
import json
import os
import re

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "handoff", "TICVAI_Readiness_Questions.xlsx")
HEAD = PatternFill("solid", fgColor="0B1324")

# The five sessions of the 28 September readiness report, by contract (and pack where a contract splits).
SESSIONS = [
    ("S1", "Access control, ticket media and credentials", "Client operations / security"),
    ("S2", "Pricing and promotions", "Client commercial / product"),
    ("S3", "Resale, orders and the approvals engine", "Client operations / finance"),
    ("S4", "Consent, waivers, privacy and customer service (with CF-127, CF-165)", "Client product / counsel"),
    ("S5", "Partner, reseller and membership terms", "Client commercial"),
]


def session_of(contract, pack):
    p = (pack or "").lower()
    if contract == "access":
        return "S1"
    if contract == "promotions":
        return "S2"
    if contract == "catalogue":
        return "S3" if "resale" in p else "S2"
    if contract in ("orders", "approvals"):
        return "S3"
    if contract == "marketing-crm":
        return "S4"
    if contract == "subscription":
        return "S5"
    return "S3"


def provisional_rows():
    """The review tool's own rows (citation, fields, flag), so both read the same way."""
    spec = importlib.util.spec_from_file_location("bpr", os.path.join(ROOT, "tools", "build-provisional-review.py"))
    bpr = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bpr)
    rows = []
    for sub in ("spine", "satellite"):
        for fn in sorted(glob.glob(os.path.join(ROOT, "contracts", sub, "*.yaml"))):
            doc = yaml.safe_load(open(fn, encoding="utf-8")) or {}
            schemas = (doc.get("components") or {}).get("schemas") or {}
            for path, ops in (doc.get("paths") or {}).items():
                for verb, op in (ops or {}).items():
                    if not isinstance(op, dict) or not op.get("x-ticvai-provisional"):
                        continue
                    desc = str(op.get("description") or "")
                    mc = bpr.CITE.search(desc)
                    schema, props = bpr.fields_of(op, schemas)
                    contract = os.path.basename(fn)[:-5]
                    pack = bpr.clean(mc.group(1)) if mc else "(uncited)"
                    rows.append({
                        "session": session_of(contract, pack), "contract": contract, "op": op["operationId"],
                        "call": f"{verb.upper()} {path}", "summary": bpr.clean(op.get("summary")),
                        "pack": pack, "page": int(mc.group(2)) if mc else 0,
                        "fields": ", ".join(props[:30]) + (" ..." if len(props) > 30 else ""),
                        "screens": "; ".join(bpr.clean(s) for s in (op.get("x-ticvai-consumed-by") or [])),
                        "flag": re.sub(r"\*\*", "", bpr.caution(props)),
                    })
    return rows


def screens():
    out = []
    for f in sorted(glob.glob(os.path.join(ROOT, "screens", "P*.yaml"))):
        doc = yaml.safe_load(open(f, encoding="utf-8")) or {}
        for s in doc.get("screens") or []:
            out.append((os.path.basename(f)[:3], s))
    return out


def sheet(wb, title, cols, widths, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    prov = provisional_rows()
    scr = screens()
    report = json.load(open(os.path.join(ROOT, "handoff", "package-report.json"), encoding="utf-8"))
    decisions = json.load(open(os.path.join(ROOT, "handoff", "audit-decisions.json"), encoding="utf-8"))
    packs = json.load(open(os.path.join(ROOT, "sources", "packs-index.json"), encoding="utf-8"))

    by_session = collections.defaultdict(list)
    for r in prov:
        by_session[r["session"]].append(r)
    prov_ops = {r["op"]: r for r in prov}

    # Per platform: operations it calls, unagreed ones by session, screens with no operation, frames.
    plat = collections.defaultdict(lambda: {"ops": set(), "noop": [], "wf": collections.Counter(), "name": ""})
    for code, s in scr:
        p = plat[code]
        ops = {a.get("operationId") for a in s.get("apis") or [] if isinstance(a, dict)} - {None}
        p["ops"] |= ops
        if not ops and not (s.get("deferred") or str(s.get("wave")) == "4"):
            p["noop"].append(s)
        p["wf"][(s.get("wireframe") or {}).get("status", "none")] += 1
    names = {p["code"]: p["name"] for p in report.get("platforms", [])}
    wb = Workbook()

    summary = []
    for code in sorted(plat):
        p = plat[code]
        unagreed = collections.Counter(prov_ops[o]["session"] for o in p["ops"] if o in prov_ops)
        pr = next((x for x in report.get("platforms", []) if x["code"] == code), {})
        closes = []
        if unagreed:
            closes.append("sessions " + ", ".join(f"{k} ({v})" for k, v in sorted(unagreed.items())))
        if p["noop"]:
            closes.append(f"{len(p['noop'])} screens to specify")
        closes.append("wireframes verified" if p["wf"].get("approved") == sum(p["wf"].values())
                      else f"wireframes: {sum(p['wf'].values()) - p['wf'].get('approved', 0)} to verify")
        summary.append((code, names.get(code, ""), pr.get("screens", ""), len(p["ops"]),
                        sum(unagreed.values()), len(p["noop"]), "; ".join(closes)))
    sheet(wb, "Summary", ["Platform", "Name", "Screens", "Operations called", "Unagreed", "Screens with no operation",
                          "What gets it to 100%"], [9, 30, 9, 12, 11, 12, 70], summary, first=True)

    sess = []
    for sid, title, who in SESSIONS:
        rs = by_session.get(sid, [])
        packs_in = collections.Counter(r["pack"] for r in rs)
        plats = sorted({s.split()[0] for r in rs for s in r["screens"].split("; ") if s})
        sess.append((sid, title, who, len(rs), sum(1 for r in rs if r["flag"]),
                     sum(1 for r in rs if not r["flag"]), ", ".join(plats),
                     "; ".join(f"{k} ({v})" for k, v in packs_in.most_common())))
    ai = [s for c, s in scr if c == "P09" and not ({a.get("operationId") for a in s.get("apis") or []
                                                    if isinstance(a, dict)} - {None})]
    sess.append(("D1", "Scoping decision: are the AI configuration assistant, forecasting and AI governance "
                 "in this delivery? If yes, they need contracts written from scratch.", "Client product",
                 0, "", "", "P09", f"{len(ai)} P09 screens name no operation"))
    sheet(wb, "Sessions", ["Id", "Session", "Who answers", "Unagreed operations", "Need a real answer",
                           "Agree-as-drafted", "Platforms", "Design packs (operations)"],
          [6, 50, 26, 11, 11, 11, 22, 70], sess)

    flagged = [(r["session"], r["pack"], r["page"] or "", r["op"], r["call"], r["summary"],
                f"Page {r['page']} of '{r['pack']}' was drafted into `{r['op']}`. {r['flag']} "
                "Which fields does this screen really record or show, and which are its filters or its behaviour?",
                r["fields"], r["screens"], "")
               for r in sorted(prov, key=lambda r: (r["session"], r["pack"], r["page"])) if r["flag"]]
    sheet(wb, "Flagged", ["Session", "Design pack", "Page", "Operation", "Call", "Screen title", "Question",
                          "Drafted fields", "Used by", "Answer"], [8, 28, 6, 26, 30, 28, 60, 50, 30, 30], flagged)

    signoff = [(r["session"], r["pack"], r["page"] or "", r["op"], r["call"], r["summary"], r["fields"],
                r["screens"], "")
               for r in sorted(prov, key=lambda r: (r["session"], r["pack"], r["page"])) if not r["flag"]]
    sheet(wb, "Sign-off", ["Session", "Design pack", "Page", "Operation", "Call", "Screen title", "Drafted fields",
                           "Used by", "Agreed / corrected (how) / not needed"],
          [8, 28, 6, 26, 30, 28, 50, 30, 34], signoff)

    und = [d for d in packs.get("undraftedDocuments") or [] if str(d).lower().endswith(".pdf")]
    sheet(wb, "Scope", ["Design pack (nothing drafted from it yet)", "Question", "Answer"], [50, 70, 30],
          [(d, "Is this pack in this delivery? If yes, we read it and draft its operations for a session; "
               "if no, it is recorded as out of scope.", "") for d in sorted(und)])

    spec = [(c, s["id"], s.get("name", ""), s.get("module", ""),
             "What does this screen show, and what does it read and change? Name its records and actions; "
             "we write the operations.", "")
            for c, s in scr if not ({a.get("operationId") for a in s.get("apis") or [] if isinstance(a, dict)} - {None})
            and not (s.get("deferred") or str(s.get("wave")) == "4")]
    sheet(wb, "Specify", ["Platform", "Screen", "Name", "Module", "Question", "Answer"], [9, 10, 36, 24, 60, 30], spec)

    wf = [(code, names.get(code, ""), sum(plat[code]["wf"].values()),
           ", ".join(f"{k} {v}" for k, v in plat[code]["wf"].most_common()),
           "Who verifies this platform's frames, and by when? Frames arrive by batch; each is signed off "
           "within 3 working days (audit R252).", "") for code in sorted(plat)]
    sheet(wb, "Wireframes", ["Platform", "Name", "Screens", "Frame status now", "Question", "Answer"],
          [9, 30, 9, 40, 60, 30], wf)

    ov = []
    for d in decisions:
        for m in re.finditer(r"\[client to (?:name|set)[^\]]*\]", json.dumps(d, ensure_ascii=False)):
            ov.append((d["id"], d["theme"], m.group(0), d["who"], ""))
    for rid, what, who in (("R126 / R205", "The age below which a guest is a minor", "Client counsel"),
                           ("R149 (3)", "How long each kind of guest document is kept", "Client counsel"),
                           ("R077 (b)", "The facial-reader vendor, which sets the SDK and template format", "Client IT")):
        ov.append((rid, what, "open", who, ""))
    sheet(wb, "Open values", ["Decision", "Topic", "Missing value", "Who names it", "Answer"], [14, 40, 50, 30, 30], ov)

    wb.save(OUT)
    print(f"{len(prov)} unagreed operations ({len(flagged)} flagged, {len(signoff)} sign-off), {len(und)} undrafted "
          f"packs, {len(spec)} screens to specify, {len(ov)} open values -> {os.path.relpath(OUT, ROOT)}")


if __name__ == "__main__":
    main()
