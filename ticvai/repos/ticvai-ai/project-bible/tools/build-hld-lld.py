# -*- coding: utf-8 -*-
"""Build the HLD and LLD pack: handoff/hld-lld/ and its Claude Design folder.

    python tools/build-hld-lld.py

Writes
    handoff/hld-lld/README.md
    handoff/hld-lld/TICVAI-HLD.md, TICVAI-HLD.svg     classic high-level design (boxes and connection types)
    handoff/hld-lld/TICVAI-LLD.md, TICVAI-LLD.svg     Azure low-level design (region, VNet, tiers, security)
    handoff/hld-lld/TICVAI - Azure Cloud Specs & Cost.xlsx
    handoff/design-batches/HLD-LLD/BRIEF.md, architecture.json

The format follows the three samples Chinmay supplied on 30 September (a classic HLD, an Azure LLD and a
specs-and-cost workbook with and without high availability). The content comes from the package: the five
deployables and their modules (handoff/service-decomposition.json, ADR-0055), the Terraform cell defaults
(repos/ticvai-infra/terraform/modules/cell), and the decisions of 30 September (ADR-0049 Qdrant one
collection per tenant, UAE hosting only; ADR-0056; ADR-0057 broker is the client's choice; ADR-0058).
Prices are estimates for Azure UAE North and say so; the workbook keeps them in editable cells.
"""
import io
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "handoff" / "hld-lld"
DESIGN = ROOT / "handoff" / "design-batches" / "HLD-LLD"
DATE = "30 September 2026"


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    with io.open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------------------------------------------------------------------------------------------------------
# The architecture, once. Both diagrams, both documents and the Claude Design brief read from here.
# ---------------------------------------------------------------------------------------------------------
dec = json.loads((ROOT / "handoff" / "service-decomposition.json").read_text(encoding="utf-8"))
MODULES = {}
for name, svc in dec["services"].items():
    MODULES.setdefault(svc.get("deployable", "?"), []).append(name.replace("Service", ""))
DEPLOYABLES = dec.get("deployables", {})

# Connection types: the sample's legend, plus the two TICVAI needs (events, local network).
LINK = {
    "api":      ("Real-time API call (HTTPS)",              "#F2A900", None),
    "direct":   ("Real-time direct call (database, cache)", "#D62828", None),
    "event":    ("Event through the broker",                "#7B2CBF", None),
    "batch":    ("Batch process",                           "#111111", "8,5"),
    "sync":     ("Near-real-time sync (offline-first)",     "#F2A900", "8,5"),
    "internet": ("Internet (TLS) / VPN",                    "#2F80ED", None),
    "wireless": ("4G / Wi-Fi",                              "#2E9E44", None),
    "lan":      ("Venue local network",                     "#2E9E44", "3,4"),
    "repl":     ("Replication",                             "#D62828", "8,5"),
}

# id: (label, sublabel, x, y, w, h). Columns: stores the deployables use (x=270) with their replicas and
# archives beside them (x=70); deployables (x=540); AI stores and the device API (x=780); front ends (x=1000);
# the outside world (x=1300). Lines between neighbouring columns run in the gaps; longer ones are routed.
HLD_NODES = {
    "replica":   ("Read replicas", "x2, zone-redundant", 70, 110, 150, 70),
    "blob":      ("Blob storage", "media, exports, backups", 70, 510, 150, 70),
    "pg":        ("PostgreSQL 16", "control DB + one DB per tenant", 270, 110, 180, 70),
    "rpt":       ("Reporting replica", "lag-tolerant, no writes", 270, 210, 180, 70),
    "redis":     ("Redis", "sessions, idempotency, cache", 270, 310, 180, 70),
    "broker":    ("Event broker", "RabbitMQ or Kafka (client)", 270, 410, 180, 70),
    "workers":   ("workers", "outbox relay, consumers, jobs", 270, 510, 180, 70),
    "commerce":  ("commerce", "sale path", 540, 110, 160, 70),
    "operations":("operations", "back office, engagement", 540, 210, 160, 70),
    "ai":        ("ticvai-ai", "AI engine (Python)", 540, 330, 160, 70),
    "access":    ("access", "gate hot path", 540, 510, 160, 70),
    "waf":       ("Front Door + WAF", "one entry, origin in UAE", 780, 60, 170, 60),
    "qdrant":    ("Qdrant", "one collection per tenant", 780, 300, 150, 70),
    "llm":       ("Azure OpenAI", "LLM, UAE region", 780, 410, 150, 60),
    "posapi":    ("Device API", "POS, kiosk, KDS", 780, 600, 160, 70),
    "fe_b2c":    ("Guest web", "B2C front end", 1000, 110, 160, 60),
    "fe_b2b":    ("Partner Portal", "B2B front end", 1000, 210, 160, 60),
    "fe_bo":     ("Back office", "Venue Mgmt, CMS, Console", 1000, 310, 160, 60),
    "pubapi":    ("Public API", "partners' systems, OTAs", 1000, 410, 160, 60),
    "u_b2c":     ("B2C guest", "web and mobile app", 1300, 110, 150, 60),
    "u_b2b":     ("B2B partner", "hotels, agents", 1300, 210, 150, 60),
    "u_ops":     ("Venue operators", "and TICVAI staff", 1300, 310, 150, 60),
    "ext3p":     ("3rd parties", "OTAs, ERP, developers", 1300, 410, 150, 60),
    "pay":       ("Payments", "Stripe, Network Intl.", 1300, 540, 150, 60),
    "einv":      ("E-invoicing", "accredited provider", 1300, 630, 150, 60),
    "msg":       ("Messaging", "email, SMS, WhatsApp, push", 1300, 720, 150, 60),
    "pos":       ("POS", "terminal, tablet", 90, 850, 120, 60),
    "kds":       ("Kitchen display", "stations", 230, 850, 120, 60),
    "kiosk":     ("Kiosks", "self-service", 370, 850, 120, 60),
    "edge":      ("Venue edge node", "access + local DB, offline", 560, 800, 170, 70),
    "turn":      ("Turnstiles", "gates, readers", 800, 850, 120, 60),
    "hand":      ("Handhelds", "scanners, staff app", 940, 850, 120, 60),
}
HLD_ZONES = [
    # (label, x, y, w, h, fill, stroke)
    ("TICVAI AZURE CLOUD - UAE North (one cell per region)", 40, 20, 1160, 680, "#FFF3CC", "#333333"),
    ("VENUE SITE (each venue)", 40, 730, 1160, 250, "#DCE6F5", "#333333"),
    ("ONSITE SALES", 70, 780, 440, 180, "#F7CFB5", "#333333"),
    ("ACCESS CONTROL", 780, 780, 300, 180, "#F7CFB5", "#333333"),
]
# (from, to, type or "type+type" for a double line, waypoints)
HLD_LINKS = [
    ("u_b2c", "fe_b2c", "internet", None), ("u_b2b", "fe_b2b", "internet", None),
    ("u_ops", "fe_bo", "internet", None), ("ext3p", "pubapi", "internet", None),
    ("fe_b2c", "waf", "api", None), ("fe_b2b", "waf", "api", None), ("fe_bo", "waf", "api", None),
    ("pubapi", "waf", "api", None),
    ("waf", "commerce", "api", None), ("waf", "operations", "api", None),
    ("commerce", "pg", "direct", None), ("commerce", "redis", "direct", None),
    ("operations", "pg", "direct", None), ("operations", "rpt", "direct", None), ("access", "pg", "direct", None),
    ("pg", "replica", "repl", None), ("pg", "rpt", "repl", None),
    ("commerce", "broker", "event", None), ("operations", "broker", "event", None), ("ai", "broker", "event", None),
    ("workers", "broker", "event", None),
    ("workers", "pg", "direct", [(252, 530), (252, 160)]),
    ("workers", "blob", "batch", None),
    ("ai", "qdrant", "direct", None), ("ai", "llm", "api", None),
    ("commerce", "pay", "api", [(745, 170), (745, 562)]),
    ("workers", "einv", "sync", [(350, 686), (1250, 686)]),
    ("workers", "msg", "sync", [(372, 693), (1240, 693)]),
    ("pos", "posapi", "internet+sync", [(150, 790)]), ("kiosk", "posapi", "internet+api", [(430, 790)]),
    ("posapi", "commerce", "api", [(760, 600), (760, 185), (700, 170)]),
    ("pos", "kds", "lan", None),
    ("edge", "access", "internet+sync", None),
    ("turn", "edge", "lan", None), ("hand", "edge", "wireless", [(1000, 815)]),
]

# LLD boxes: id: (label, lines, x, y, w, h, fill)
LLD_NODES = {
    "users":   ("Users", ["B2C guests", "B2B partners", "Venue operators"], 30, 330, 150, 110, "#FFFFFF"),
    "fd":      ("Azure Front Door Premium", ["WAF: OWASP + bot rules", "TLS 1.2+, custom domains", "Private Link to origin"], 220, 320, 190, 130, "#FFFFFF"),
    "ilb":     ("Ingress (internal LB)", ["NGINX ingress on AKS", "Private Link service"], 470, 330, 170, 110, "#EEF4FF"),
    "sys":     ("AKS system pool", ["D4s v5, 2-4 nodes", "zones 1-3"], 690, 150, 190, 80, "#EEF4FF"),
    "wl":      ("AKS workload pool", ["D8s v5, 3-20 nodes", "commerce, access,", "operations, workers"], 690, 250, 190, 110, "#EEF4FF"),
    "aip":     ("AKS AI pool", ["D8s v5, 2-4 nodes", "ticvai-ai, embeddings,", "reranker (CPU)"], 690, 380, 190, 110, "#EEF4FF"),
    "data":    ("AKS data pool", ["E4s v5 x3, zones 1-3", "Qdrant 3-node cluster", "broker 3-node cluster"], 690, 510, 190, 110, "#EEF4FF"),
    "pg":      ("PostgreSQL Flexible Server 16", ["GP D4ds v5, zone-redundant HA", "control DB + DB per tenant", "PITR 35 days, UAE only"], 950, 150, 220, 110, "#FFFFFF"),
    "ro":      ("Read replicas x2", ["GP D4ds v5", "gate checks never read here"], 950, 280, 220, 70, "#FFFFFF"),
    "rpt":     ("Reporting + AI log DB", ["reporting replica (D2ds v5)", "AI log DB (D4ds v5, 1 TB)"], 950, 370, 220, 80, "#FFFFFF"),
    "pe":      ("Private endpoints", ["Redis Premium (13 GB)", "Key Vault (HSM keys)", "Blob storage (ZRS)", "Container Registry"], 950, 470, 220, 120, "#FFFFFF"),
    "bastion": ("Azure Bastion", ["admin access only", "no public SSH / RDP"], 470, 560, 170, 80, "#FFFFFF"),
    "nat":     ("NAT Gateway", ["one static egress IP", "for provider allow-lists"], 470, 700, 170, 80, "#FFFFFF"),
    "extp":    ("External providers", ["Payments, e-invoicing", "messaging, Azure OpenAI"], 220, 700, 190, 80, "#FFFFFF"),
    "ops":     ("Operations", ["Azure Monitor + Log Analytics", "Azure Backup, Defender", "Entra ID, subscription"], 30, 560, 180, 100, "#FFFFFF"),
    "venue":   ("Venue site", ["POS, kiosks, KDS", "venue edge node (offline)", "turnstiles, handhelds on LAN"], 1250, 300, 190, 110, "#FFFFFF"),
    "vpn":     ("Site-to-site VPN", ["dedicated tier only;", "others: Internet + mTLS"], 1250, 450, 190, 80, "#FFFFFF"),
}
# Frames a line may start or end on without being a box: (label, x, y, w, h)
LLD_FRAMES = {"aks": ("AKS cluster", 670, 130, 230, 510)}
LLD_LINKS = [
    ("users", "fd", "internet", None), ("fd", "ilb", "api", None),
    ("ilb", "wl", "api", None), ("ilb", "aip", "api", None),
    ("wl", "pg", "direct", None), ("wl", "ro", "direct", None), ("wl", "pe", "direct", None),
    ("aip", "rpt", "direct", None), ("aip", "data", "direct", None),
    ("wl", "data", "event", [(912, 345), (912, 580)]),
    ("pg", "ro", "repl", None), ("pg", "rpt", "repl", [(1180, 215), (1180, 405)]),
    ("aks", "nat", "api", [(785, 740)]), ("nat", "extp", "api", None),
    ("bastion", "sys", "internet", [(660, 600), (660, 190)]),
    ("venue", "fd", "internet", [(1345, 85), (315, 85)]),
    ("venue", "vpn", "internet", None),
    ("vpn", "ilb", "sync", [(1225, 490), (1225, 676), (650, 676), (650, 420)]),
]
SUBNETS = [
    # (name, cidr, holds, rule)
    ("snet-ingress", "10.20.0.0/24", "Internal load balancer and Private Link service for Front Door",
     "Inbound only from Front Door's Private Link; no public IP"),
    ("snet-aks-system", "10.20.4.0/22", "AKS system node pool", "No inbound from outside the VNet"),
    ("snet-aks-workload", "10.20.8.0/21", "Workload pool: commerce, access, operations, workers",
     "Inbound from snet-ingress only; outbound through the NAT Gateway"),
    ("snet-aks-ai", "10.20.16.0/22", "AI pool: ticvai-ai, embeddings, reranker",
     "Inbound from snet-ingress and snet-aks-workload; read-only role on transactional schemas (ADR-0020)"),
    ("snet-aks-data", "10.20.20.0/23", "Data pool: Qdrant cluster, broker cluster",
     "Inbound from the workload and AI pools only; Qdrant needs each tenant's collection-scoped JWT"),
    ("snet-postgres", "10.20.24.0/24", "PostgreSQL Flexible Server (delegated subnet), replicas, AI log DB",
     "Inbound 5432 from the AKS subnets only (through pgbouncer); public access disabled"),
    ("snet-private-endpoints", "10.20.25.0/24", "Redis, Key Vault, Blob storage, Container Registry",
     "Private endpoints with private DNS zones; public access disabled on every resource"),
    ("AzureBastionSubnet", "10.20.26.0/26", "Azure Bastion", "The only administrative path in"),
]
AKS_SERVICE_CIDR = "10.100.0.0/16"  # Terraform default (modules/cell/variables.tf)


# ---------------------------------------------------------------------------------------------------------
# SVG
# ---------------------------------------------------------------------------------------------------------
def box3d(x, y, w, h, label, sub=None):
    d = 10  # depth of the 3D face, as in the sample
    s = []
    s.append(f'<polygon points="{x},{y} {x+d},{y-d} {x+w+d},{y-d} {x+w},{y}" fill="#E6E6E6" stroke="#8A8A8A"/>')
    s.append(f'<polygon points="{x+w},{y} {x+w+d},{y-d} {x+w+d},{y+h-d} {x+w},{y+h}" fill="#9E9E9E" stroke="#8A8A8A"/>')
    s.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="url(#g)" stroke="#8A8A8A"/>')
    cy = y + h / 2 - (7 if sub else 0)
    s.append(f'<text x="{x+w/2}" y="{cy+5}" text-anchor="middle" font-size="15" font-weight="600">{esc(label)}</text>')
    if sub:
        s.append(f'<text x="{x+w/2}" y="{cy+23}" text-anchor="middle" font-size="11" fill="#333">{esc(sub)}</text>')
    return "\n".join(s)


def toward(n, pt):
    """Where the line from box n's centre towards pt leaves the box."""
    x, y, w, h = n
    cx, cy = x + w / 2, y + h / 2
    dx, dy = pt[0] - cx, pt[1] - cy
    if dx == 0 and dy == 0:
        return cx, cy
    sx = (w / 2) / abs(dx) if dx else float("inf")
    sy = (h / 2) / abs(dy) if dy else float("inf")
    t = min(sx, sy)
    return cx + dx * t, cy + dy * t


def facing(a, b):
    """Anchors on the edges of a and b that face each other, so a line between neighbours stays in the gap."""
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    gapx = max(bx - (ax + aw), ax - (bx + bw))
    gapy = max(by - (ay + ah), ay - (by + bh))
    if gapx >= 0:
        x1 = ax + aw if bx > ax else ax
        x2 = bx if bx > ax else bx + bw
        return (x1, ay + ah / 2), (x2, by + bh / 2)
    if gapy >= 0:
        y1 = ay + ah if by > ay else ay
        y2 = by if by > ay else by + bh
        return (ax + aw / 2, y1), (bx + bw / 2, y2)
    return toward(a, (bx + bw / 2, by + bh / 2)), toward(b, (ax + aw / 2, ay + ah / 2))


def route(a, b, via):
    if not via:
        return list(facing(a, b))
    return [toward(a, via[0])] + list(via) + [toward(b, via[-1])]


def polyline(points, kind, off=(0, 0)):
    _, color, dash = LINK[kind]
    da = f' stroke-dasharray="{dash}"' if dash else ""
    pts = " ".join(f"{x + off[0]:.0f},{y + off[1]:.0f}" for x, y in points)
    return f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="2.4"{da}/>'


def draw_link(points, kinds):
    kinds = kinds.split("+")
    if len(kinds) == 1:
        return polyline(points, kinds[0])
    (x1, y1), (x2, y2) = points[0], points[-1]
    off = (0, 2.6) if abs(x2 - x1) >= abs(y2 - y1) else (2.6, 0)
    return "\n".join([polyline(points, kinds[0], (-off[0], -off[1])), polyline(points, kinds[1], off)])


def seg_hits(x1, y1, x2, y2, l, t, r, bt):
    p = [-(x2 - x1), x2 - x1, -(y2 - y1), y2 - y1]
    q = [x1 - l, r - x1, y1 - t, bt - y1]
    u1, u2 = 0.0, 1.0
    for pi, qi in zip(p, q):
        if pi == 0:
            if qi < 0:
                return False
            continue
        u = qi / pi
        if pi < 0:
            u1 = max(u1, u)
        else:
            u2 = min(u2, u)
    return u1 < u2


def crossings(boxes, links):
    """Every line segment that runs through a box other than its own two ends. Must be empty."""
    bad = []
    for a, b, kinds, pts in links:
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            for k, (x, y, w, h) in boxes.items():
                if k in (a, b):
                    continue
                if seg_hits(x1, y1, x2, y2, x + 3, y + 3, x + w - 3, y + h - 3):
                    bad.append(f"{a} -> {b} crosses {k}")
    return sorted(set(bad))


def legend(x, y, kinds):
    h = 50 + 26 * len(kinds)
    s = [f'<rect x="{x}" y="{y}" width="330" height="{h}" fill="#FFFFFF" stroke="#111" stroke-width="2"/>',
         f'<text x="{x+30}" y="{y+34}" font-size="20" font-weight="700">CONNECTIONS</text>']
    for i, k in enumerate(kinds):
        label, color, dash = LINK[k]
        yy = y + 60 + 26 * i
        da = f' stroke-dasharray="{dash}"' if dash else ""
        s.append(f'<line x1="{x+18}" y1="{yy}" x2="{x+78}" y2="{yy}" stroke="{color}" stroke-width="2.4"{da}/>')
        s.append(f'<text x="{x+90}" y="{yy+5}" font-size="13">{esc(label)}</text>')
    return "\n".join(s)


SVG_HEAD = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {w} {h}" width="{w}" height="{h}" '
            'font-family="Segoe UI, Arial, sans-serif" fill="#111">\n'
            '<defs><linearGradient id="g" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#D9D9D9"/>'
            '<stop offset="1" stop-color="#BDBDBD"/></linearGradient></defs>\n'
            '<rect width="100%" height="100%" fill="#FFFFFF"/>\n')


CROSS = {}


def hld_svg():
    s = [SVG_HEAD.format(w=1570, h=1150)]
    for label, x, y, w, h, fill, stroke in HLD_ZONES:
        s.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{stroke}" stroke-width="1.2"/>')
        big = label.startswith("TICVAI") or label.startswith("VENUE SITE")
        ly = y + 28 if big else y + h - 16
        s.append(f'<text x="{x+14}" y="{ly}" font-size="{20 if big else 18}">{esc(label)}</text>')
    geo = {k: (v[2], v[3], v[4], v[5]) for k, v in HLD_NODES.items()}
    drawn = [(a, b, kind, route(geo[a], geo[b], via)) for a, b, kind, via in HLD_LINKS]
    CROSS["hld"] = crossings(geo, drawn)
    for a, b, kind, pts in drawn:
        s.append(draw_link(pts, kind))
    for k, (label, sub, x, y, w, h) in HLD_NODES.items():
        s.append(box3d(x, y, w, h, label, sub))
    s.append('<text x="1300" y="36" font-size="24" font-weight="700">TICVAI PLATFORM HLD</text>')
    s.append(f'<text x="1300" y="58" font-size="12" fill="#555">{DATE}</text>')
    s.append(legend(1220, 830, ["api", "direct", "event", "batch", "sync", "internet", "wireless", "lan", "repl"]))
    s.append("</svg>\n")
    return "\n".join(s)


def lld_box(x, y, w, h, title, lines, fill):
    s = [f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="8" fill="{fill}" stroke="#2F6DB5" stroke-width="1.4"/>',
         f'<text x="{x+w/2}" y="{y+22}" text-anchor="middle" font-size="13.5" font-weight="700">{esc(title)}</text>']
    for i, t in enumerate(lines):
        s.append(f'<text x="{x+w/2}" y="{y+42+i*17}" text-anchor="middle" font-size="11.5" fill="#333">{esc(t)}</text>')
    return "\n".join(s)


def lld_svg():
    s = [SVG_HEAD.format(w=1600, h=820)]
    s.append('<text x="30" y="40" font-size="24" font-weight="700">TICVAI PLATFORM LLD - Azure</text>')
    s.append(f'<text x="30" y="64" font-size="12" fill="#555">{DATE} - production cell, with high availability</text>')
    s.append('<rect x="450" y="100" width="760" height="560" rx="6" fill="#F4FBF4" stroke="#2E9E44" stroke-width="2.4"/>')
    s.append('<text x="466" y="126" font-size="16" font-weight="700">Region "UAE North" - VNet 10.20.0.0/16</text>')
    s.append('<text x="466" y="654" font-size="11.5" fill="#333">zones 1-3; every store in the UAE (backups and DR included)</text>')
    for label, x, y, w, h in [("AKS cluster (Standard tier)", 670, 130, 230, 510), ("Data tier", 935, 130, 250, 480)]:
        s.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="6" fill="none" stroke="#6A7FA8" stroke-dasharray="6,4"/>')
        s.append(f'<text x="{x+10}" y="{y+h-8}" font-size="12" fill="#2F4B7C">{esc(label)}</text>')
    geo = {k: (v[2], v[3], v[4], v[5]) for k, v in LLD_NODES.items()}
    ends = dict(geo, **{k: v[1:] for k, v in LLD_FRAMES.items()})
    drawn = [(a, b, kind, route(ends[a], ends[b], via)) for a, b, kind, via in LLD_LINKS]
    CROSS["lld"] = crossings(geo, drawn)
    for a, b, kind, pts in drawn:
        s.append(draw_link(pts, kind))
    for k, (title, lines, x, y, w, h, fill) in LLD_NODES.items():
        s.append(lld_box(x, y, w, h, title, lines, fill))
    s.append(legend(1250, 580, ["internet", "api", "direct", "event", "sync", "repl"]))
    s.append("</svg>\n")
    return "\n".join(s)


# ---------------------------------------------------------------------------------------------------------
# Cost workbook. Unit prices are USD per month (730 hours, Linux, pay as you go) estimated for UAE North;
# every one is an editable cell, and the totals are formulas.
# ---------------------------------------------------------------------------------------------------------
AED = 3.6725
MARGIN = 0.20
# (group, service type, description, specs, qty, unit USD, scaled qty or None = same)
PROD_HA = [
    ("Compute (AKS)", "Azure Kubernetes Service", "Control plane, Standard tier",
     "Uptime SLA, zone-redundant control plane", 1, 73, None),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "System node pool",
     "D4s v5 (4 vCPU, 16 GB), Linux, zones 1-3", 3, 175, None),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "Workload pool: commerce, access, operations, workers",
     "D8s v5 (8 vCPU, 32 GB), Linux, autoscale 3-20 across zones", 3, 350, 6),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "AI pool: ticvai-ai, embeddings, reranker (CPU)",
     "D8s v5 (8 vCPU, 32 GB), Linux, autoscale 2-4", 2, 350, 4),
    ("Vector store", "Virtual Machines (AKS nodes)", "Qdrant cluster, one collection per tenant (ADR-0049)",
     "E4s v5 (4 vCPU, 32 GB) x3, zones 1-3, replication factor 2", 3, 225, None),
    ("Vector store", "Managed Disks", "Qdrant storage", "Premium SSD P15 256 GB per node", 3, 38, None),
    ("Event broker", "Virtual Machines (AKS nodes)", "Broker cluster: RabbitMQ or Kafka (client's choice, ADR-0057)",
     "D2s v5 (2 vCPU, 8 GB) x3, zones 1-3; re-price if a managed service is chosen", 3, 88, None),
    ("Database", "Azure Database for PostgreSQL", "Primary with zone-redundant standby: control DB + DB per tenant",
     "Flexible Server 16, General Purpose D4ds v5 (4 vCore, 16 GB) x2 (primary + standby)", 2, 320, None),
    ("Database", "Azure Database for PostgreSQL", "Storage, primary and standby",
     "256 GB Premium SSD each; PITR backup 35 days included up to storage size", 2, 36, None),
    ("Database", "Azure Database for PostgreSQL", "Read replicas (reads; gate checks never read here)",
     "General Purpose D4ds v5 + 256 GB each", 2, 356, None),
    ("Database", "Azure Database for PostgreSQL", "Reporting replica (lag-tolerant, no writes)",
     "General Purpose D2ds v5 (2 vCore, 8 GB) + 256 GB", 1, 196, None),
    ("Database", "Azure Database for PostgreSQL", "AI log database (decision records, prompts)",
     "General Purpose D4ds v5 + 1 TB", 1, 460, None),
    ("Cache", "Azure Cache for Redis", "Sessions, idempotency keys, caches, AI features",
     "Premium P2 (13 GB), zone-redundant, private endpoint", 1, 1000, None),
    ("Storage", "Storage account (Blob, ZRS)", "Media, exports, Qdrant snapshots, database dumps",
     "Hot tier, zone-redundant, 2 TB, UAE North only", 1, 60, None),
    ("Storage", "Container Registry", "Images for the five deployables", "Standard", 1, 20, None),
    ("Security", "Key Vault", "Secrets, per-tenant keys, Qdrant JWT signing keys", "Premium (HSM-backed keys)", 1, 10, None),
    ("Network", "Azure Front Door", "Single entry for web, apps and APIs, with WAF",
     "Premium: WAF managed rules and bot protection, Private Link origin; base + estimated traffic", 1, 450, 2),
    ("Network", "NAT Gateway", "One static egress IP for payment and e-invoicing allow-lists",
     "730 hours, 1 TB processed", 1, 80, None),
    ("Network", "Azure Bastion", "Administrative access only", "Basic, 730 hours", 1, 139, None),
    ("Network", "Private Link", "Private endpoints and private DNS zones", "About 6 endpoints", 1, 50, None),
    ("Network", "Bandwidth", "Outbound data transfer", "About 2 TB a month", 1, 180, 2),
    ("Operations", "Azure Monitor / Log Analytics", "Logs, metrics, alerts, dashboards", "About 60 GB ingested a month", 1, 200, None),
    ("Operations", "Azure Backup", "Snapshots of Qdrant and broker disks", "Daily, 30 days, UAE North", 1, 40, None),
    ("Operations", "Microsoft Defender for Cloud", "Servers, containers and databases", "Estimated", 1, 150, None),
    ("AI (usage)", "Azure OpenAI", "Language model calls (re-billed per token, AI-D02)",
     "Usage-based; not a fixed hosting cost", 1, 0, None),
]
PROD_NO_HA = [
    ("Compute (AKS)", "Azure Kubernetes Service", "Control plane, Standard tier", "Uptime SLA", 1, 73, None),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "System node pool", "D4s v5 (4 vCPU, 16 GB), Linux, one zone", 2, 175, None),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "Workload pool: commerce, access, operations, workers",
     "D8s v5 (8 vCPU, 32 GB), Linux, autoscale 2-10", 2, 350, 4),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "AI pool: ticvai-ai, embeddings, reranker (CPU)",
     "D8s v5 (8 vCPU, 32 GB), Linux, autoscale 1-2", 1, 350, 2),
    ("Vector store", "Virtual Machines (AKS nodes)", "Qdrant, one collection per tenant (ADR-0049)",
     "E4s v5 (4 vCPU, 32 GB), single node", 1, 225, None),
    ("Vector store", "Managed Disks", "Qdrant storage", "Premium SSD P15 256 GB", 1, 38, None),
    ("Event broker", "Virtual Machines (AKS nodes)", "Broker: RabbitMQ or Kafka (client's choice, ADR-0057)",
     "D2s v5 (2 vCPU, 8 GB), single node", 1, 88, None),
    ("Database", "Azure Database for PostgreSQL", "Primary: control DB + DB per tenant",
     "Flexible Server 16, General Purpose D4ds v5 (4 vCore, 16 GB), no standby", 1, 320, None),
    ("Database", "Azure Database for PostgreSQL", "Storage", "256 GB Premium SSD; PITR backup 35 days", 1, 36, None),
    ("Database", "Azure Database for PostgreSQL", "Read replica (also serves reporting)",
     "General Purpose D4ds v5 + 256 GB", 1, 356, None),
    ("Database", "Azure Database for PostgreSQL", "AI log database", "General Purpose D2ds v5 + 512 GB", 1, 230, None),
    ("Cache", "Azure Cache for Redis", "Sessions, idempotency keys, caches, AI features", "Standard C4 (13 GB)", 1, 400, None),
    ("Storage", "Storage account (Blob, LRS)", "Media, exports, Qdrant snapshots, database dumps", "Hot tier, 2 TB, UAE North only", 1, 45, None),
    ("Storage", "Container Registry", "Images for the five deployables", "Standard", 1, 20, None),
    ("Security", "Key Vault", "Secrets, per-tenant keys, Qdrant JWT signing keys", "Standard", 1, 5, None),
    ("Network", "Azure Front Door", "Single entry for web, apps and APIs, with WAF",
     "Premium: WAF managed rules and bot protection; base + estimated traffic", 1, 450, 2),
    ("Network", "NAT Gateway", "One static egress IP for payment and e-invoicing allow-lists", "730 hours, 1 TB processed", 1, 80, None),
    ("Network", "Azure Bastion", "Administrative access only", "Basic, 730 hours", 1, 139, None),
    ("Network", "Bandwidth", "Outbound data transfer", "About 2 TB a month", 1, 180, 2),
    ("Operations", "Azure Monitor / Log Analytics", "Logs, metrics, alerts", "About 40 GB ingested a month", 1, 140, None),
    ("Operations", "Microsoft Defender for Cloud", "Servers, containers and databases", "Estimated", 1, 100, None),
    ("AI (usage)", "Azure OpenAI", "Language model calls (re-billed per token, AI-D02)", "Usage-based", 1, 0, None),
]
PREPROD = [
    ("Compute (AKS)", "Azure Kubernetes Service", "Control plane", "Free tier", 1, 0, None),
    ("Compute (AKS)", "Virtual Machines (AKS nodes)", "System and workload nodes, all five deployables",
     "D8s v5 (8 vCPU, 32 GB) x2 + D4s v5 x1, Linux", 1, 875, None),
    ("Vector store", "Virtual Machines (AKS nodes)", "Qdrant, single node", "E2s v5 (2 vCPU, 16 GB) + P10 128 GB", 1, 132, None),
    ("Event broker", "Virtual Machines (AKS nodes)", "Broker, single node", "D2s v5 (2 vCPU, 8 GB)", 1, 88, None),
    ("Database", "Azure Database for PostgreSQL", "Control DB + test tenant DBs + AI log DB",
     "Flexible Server 16, General Purpose D2ds v5 + 256 GB, no standby, no replicas", 1, 196, None),
    ("Cache", "Azure Cache for Redis", "Test cache", "Standard C1 (1 GB)", 1, 125, None),
    ("Storage", "Storage account (Blob, LRS)", "Test media and snapshots", "Hot tier, 500 GB", 1, 15, None),
    ("Security", "Key Vault", "Test secrets", "Standard", 1, 5, None),
    ("Network", "Front Door, NAT, Bastion, Registry", "Shared with production", "Separate routes and WAF policy", 1, 0, None),
    ("Operations", "Azure Monitor / Log Analytics", "Logs and metrics", "About 20 GB ingested a month", 1, 70, None),
]
NOTES_PROD = [
    "Prices are estimates for Azure region UAE North, pay as you go, 730 hours a month, 30 September 2026. "
    "Confirm each line in the Azure pricing calculator for UAE North before quoting; reserved instances "
    "(1 or 3 years) lower compute by roughly 30-55%.",
    "Every store is hosted in the UAE, backups and disaster recovery included. Qdrant was approved on that "
    "condition (30 September).",
    "One production cell serves many tenants: each tenant has its own PostgreSQL database and its own Qdrant "
    "collection with a collection-scoped key.",
    "\"When scaled up\" doubles the autoscaled lines (workload and AI pools, Front Door traffic, bandwidth) for "
    "busy periods such as an on-sale or a holiday peak.",
    "Not included: Azure OpenAI tokens (billed per use), SMS, email and WhatsApp fees, payment provider fees, "
    "the e-invoicing provider, Apple and Google developer accounts, and venue hardware.",
    "The venue edge node (one per venue, runs the gates and POS offline) is venue hardware, not Azure: "
    "recommended 8 cores, 32 GB RAM, 1 TB SSD, two units per venue for failover.",
    "The event broker is RabbitMQ or Kafka, the client's choice. It is priced self-hosted on AKS; a managed "
    "service is re-priced when chosen.",
]
NOTES_NO_HA = [
    "Single zone throughout: a zone outage stops the cloud services until they are restored. Venue gates and "
    "POS keep working on the venue edge node, offline.",
    "Not suitable for the cloud availability target (99.95% for commerce); use the high-availability sheet.",
]
NOTES_PREPROD = [
    "One pre-production environment for acceptance testing and client demos; development runs locally on "
    "docker-compose and on this cluster.",
    "Zero-priced lines are covered by production.",
]

thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3056")
ENV_FILL = PatternFill("solid", fgColor="DCE6F5")
TOT_FILL = PatternFill("solid", fgColor="FFF3CC")
WRAP = Alignment(wrap_text=True, vertical="top")


def cost_block(ws, r, env, rows, notes, scaled_col=True):
    heads = ["Environment", "Group", "Azure service type", "Service description", "Region", "Specs description",
             "Qty", "Unit price (USD / month)", "Estimated monthly cost (USD)"]
    if scaled_col:
        heads.append("When solution is scaled up (USD)")
    for c, h in enumerate(heads, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    first = r + 1
    for i, (grp, typ, desc, spec, qty, unit, sq) in enumerate(rows):
        rr = first + i
        vals = [env if i == 0 else None, grp, typ, desc, "UAE North", spec, qty, unit]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(rr, c, v)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.cell(rr, 9, f"=G{rr}*H{rr}").number_format = "#,##0"
        ws.cell(rr, 9).border = BORDER
        ws.cell(rr, 8).number_format = "#,##0"
        if scaled_col:
            ws.cell(rr, 10, f"={sq}*H{rr}" if sq else f"=I{rr}").number_format = "#,##0"
            ws.cell(rr, 10).border = BORDER
    last = first + len(rows) - 1
    ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
    envc = ws.cell(first, 1)
    envc.fill = ENV_FILL
    envc.font = Font(bold=True, size=12)
    envc.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    r = last + 1
    totals = [("USD - Total cost per month", f"=SUM(I{first}:I{last})", f"=SUM(J{first}:J{last})"),
              ("AED - Total cost per month", f"=I{r}*$C$3", f"=J{r}*$C$3"),
              ("AED - Total cost per year", f"=I{r+1}*12", None),
              ("AED - Total price per year with margin", f"=I{r+2}*(1+$C$4)", None)]
    for j, (label, f1, f2) in enumerate(totals):
        ws.cell(r + j, 6, label).font = Font(bold=True)
        ws.cell(r + j, 6).alignment = Alignment(horizontal="right")
        c = ws.cell(r + j, 9, f1)
        c.number_format = "#,##0"
        c.font = Font(bold=True)
        c.fill = TOT_FILL
        if scaled_col and f2:
            c2 = ws.cell(r + j, 10, f2)
            c2.number_format = "#,##0"
            c2.fill = TOT_FILL
    r += len(totals) + 1
    ws.cell(r, 2, f"Notes - {env}").font = Font(bold=True)
    for n in notes:
        r += 1
        ws.cell(r, 2, n).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 32
    return r + 3


def cost_sheet(wb, title, prod_rows, prod_notes):
    ws = wb.create_sheet(title)
    ws["A1"] = f"TICVAI - Azure cloud specs and cost ({title.lower()})"
    ws["A1"].font = Font(bold=True, size=14)
    ws["B3"], ws["C3"] = "AED per USD", AED
    ws["B4"], ws["C4"] = "Margin", MARGIN
    ws["C4"].number_format = "0%"
    ws["D3"] = "Edit the rate, the margin and any unit price; every total is a formula."
    r = cost_block(ws, 6, "Production", prod_rows, NOTES_PROD + prod_notes)
    cost_block(ws, r, "Pre-production", PREPROD, NOTES_PREPROD)
    widths = [16, 16, 26, 38, 11, 46, 6, 14, 16, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A7"
    return ws


def build_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    cost_sheet(wb, "Infra without HA", PROD_NO_HA, NOTES_NO_HA)
    cost_sheet(wb, "Infra with High Availability", PROD_HA, [
        "Zones 1-3 in UAE North: PostgreSQL zone-redundant standby, Qdrant and broker three-node clusters, AKS "
        "node pools across zones, Redis Premium zone-redundant."])
    ws = wb.create_sheet("Deployables")
    ws.append(["Deployable", "What it is", "Modules", "Runs on"])
    runs = {"commerce": "Workload pool", "access": "Workload pool (and the venue edge node)",
            "operations": "Workload pool", "ticvai-ai": "AI pool", "workers": "Workload pool"}
    for k, v in DEPLOYABLES.items():
        ws.append([k, v, ", ".join(sorted(MODULES.get(k, []))) or "-", runs.get(k, "")])
    ws.append([])
    ws.append(["Data stores", "", "", ""])
    for row in [("PostgreSQL 16", "Control DB and one database per tenant; read replicas; reporting replica; AI log DB", "", "Flexible Server"),
                ("Qdrant", "Vectors; one collection per tenant, collection-scoped JWT (ADR-0049)", "", "Data pool"),
                ("Redis", "Sessions, idempotency, caches", "", "Azure Cache for Redis"),
                ("Event broker", "RabbitMQ or Kafka, the client's choice (ADR-0057); outbox relay per region (ADR-0058)", "", "Data pool"),
                ("Blob storage", "Media, exports, snapshots, dumps", "", "Storage account")]:
        ws.append(list(row))
    for i, w in enumerate([18, 70, 60, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = WRAP
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    ws = wb.create_sheet("Network")
    ws.append(["Subnet", "Address range", "Holds", "Rule"])
    for sn in SUBNETS:
        ws.append(list(sn))
    ws.append(["AKS service range", AKS_SERVICE_CIDR, "Kubernetes services (Terraform default)", "Internal only"])
    for i, w in enumerate([24, 16, 60, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = WRAP
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    wb.calculation.fullCalcOnLoad = True  # the totals are formulas; Excel computes them on open
    wb.save(path)


# ---------------------------------------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------------------------------------
def hld_md():
    rows = "\n".join(f"| **{k}** | {v} | {', '.join(sorted(MODULES.get(k, []))) or '-'} |" for k, v in DEPLOYABLES.items())
    links = "\n".join(f"| {HLD_NODES[a][0]} | {HLD_NODES[b][0]} | {' and '.join(LINK[k][0] for k in t.split('+'))} |"
                      for a, b, t, _ in HLD_LINKS)
    return f"""# TICVAI platform: high-level design

> **Date:** {DATE} · **Generated by** `tools/build-hld-lld.py` · **Diagram:** [`TICVAI-HLD.svg`](TICVAI-HLD.svg)
> **Follows** the classic HLD format supplied on 30 September: the cloud, the venue site and the outside world as
> boxes, and every line typed by how the two ends talk.

![TICVAI HLD](TICVAI-HLD.svg)

## The shape in one paragraph

One cell per region, in Azure UAE North. Inside it, 17 modules in one .NET solution run as five deployables
(ADR-0055). Each tenant has its own PostgreSQL database and its own Qdrant collection (ADR-0038, ADR-0049).
Modules in different deployables talk through events on the broker, written first to an outbox in the same
transaction (ADR-0033, ADR-0058). At each venue an edge node runs the gates and the POS on its own when the
Internet is down (ADR-0013), and syncs when it is back.

## Deployables

| Deployable | What it is | Modules |
|---|---|---|
{rows}

## Data stores

| Store | Holds | Decision |
|---|---|---|
| PostgreSQL 16 | A control database, one database per tenant, read replicas, a reporting replica and the AI log database | ADR-0038, ADR-0056 (UUIDv7 ids, monthly partitions) |
| Qdrant | The AI knowledge index: one collection per tenant, each with its own collection-scoped key; venue scope is a filter the retrieval client always adds | ADR-0049 (approved 30 September on condition of UAE hosting) |
| Redis | Sessions, idempotency keys, caches | ADR-0032 |
| Event broker | Events between deployables; RabbitMQ or Kafka, the client's choice | ADR-0057, ADR-0058 |
| Blob storage | Media, exports, Qdrant snapshots, database dumps | ADR-0047 |

## Connections

| From | To | How |
|---|---|---|
{links}

## At the venue

- **Onsite sales:** POS terminals and kiosks call the device API over the Internet; the kitchen display takes
  orders from the POS on the venue network. Offline, the POS keeps selling and syncs later.
- **Access control:** turnstiles and handhelds talk to the venue edge node on the venue network or Wi-Fi. The
  edge node holds the offline package of valid entitlements and decides in under 300 ms; it syncs scans to
  the cloud as soon as it can.

## What this does not show

Per-screen flows (see `diagrams/lld/`), the AI internals (`docs/architecture/ai-system-design.md`) and sizing
(`handoff/sizing.json`). The Azure network, security and cost are in [`TICVAI-LLD.md`](TICVAI-LLD.md) and the
cost workbook.
"""


def lld_md():
    sn = "\n".join(f"| `{a}` | {b} | {c} | {d} |" for a, b, c, d in SUBNETS)
    return f"""# TICVAI platform: low-level design (Azure)

> **Date:** {DATE} · **Generated by** `tools/build-hld-lld.py` · **Diagram:** [`TICVAI-LLD.svg`](TICVAI-LLD.svg)
> **Costs:** [`TICVAI - Azure Cloud Specs & Cost.xlsx`](TICVAI%20-%20Azure%20Cloud%20Specs%20%26%20Cost.xlsx) (with and without high availability)
> **Source of sizes:** the Terraform cell module (`repos/ticvai-infra/terraform/modules/cell`), `handoff/sizing.json`
> and the decisions of 30 September.

![TICVAI LLD](TICVAI-LLD.svg)

## Region and residency

- **Region:** Azure UAE North, availability zones 1-3. **Every store stays in the UAE**, including backups,
  snapshots and disaster recovery: the condition Qdrant was approved on (30 September) and ADR-0009.
- **One cell per region.** A cell serves many tenants; each tenant has its own PostgreSQL database and its own
  Qdrant collection.

## Tiers

| Tier | What runs there | Size (production, with HA) |
|---|---|---|
| Edge | Azure Front Door Premium with WAF (OWASP and bot rules), TLS 1.2+, Private Link to the origin | One profile |
| Ingress | NGINX ingress behind an internal load balancer, published to Front Door by Private Link | In the system pool |
| AKS system pool | Kubernetes system services | D4s v5, 2-4 nodes, zones 1-3 |
| AKS workload pool | commerce, access, operations, workers | D8s v5, 3-20 nodes, autoscaled |
| AKS AI pool | ticvai-ai, the embedding model and the reranker (CPU) | D8s v5, 2-4 nodes |
| AKS data pool | Qdrant (3 nodes), the event broker (3 nodes) | E4s v5 x3 and D2s v5 x3, zones 1-3 |
| Database | PostgreSQL Flexible Server 16: primary with zone-redundant standby, 2 read replicas, a reporting replica, the AI log database | General Purpose D4ds v5; 256 GB each; AI log 1 TB |
| Private endpoints | Redis Premium (13 GB), Key Vault (HSM keys), Blob storage (ZRS), Container Registry | Public access disabled |

## Network

VNet `10.20.0.0/16` (proposed; confirm against the client's address plan before the first deploy).

| Subnet | Range | Holds | Rule |
|---|---|---|---|
{sn}
| AKS services | `{AKS_SERVICE_CIDR}` | Kubernetes service range (Terraform default) | Internal only |

- **Inbound:** only through Front Door (web, apps, APIs, venue devices) and Bastion (administrators). No
  resource has a public database, cache or storage endpoint.
- **Outbound:** through the NAT Gateway's one static IP, which payment and e-invoicing providers can allow-list.
- **Venues:** devices reach the cloud over the Internet with TLS and device certificates; a dedicated-tier
  tenant may add a site-to-site VPN. Turnstiles and handhelds only talk to the venue edge node.

## Security

- **Identity:** Entra ID for administrators; the platform's own identity for staff, guests and partners.
  Workloads use managed identities to reach Key Vault, storage and the registry.
- **Tenant isolation:** each tenant's database (with row-level security for venue scope) and each tenant's
  Qdrant collection, readable only with that tenant's collection-scoped JWT, signed with a key in Key Vault.
- **Least privilege between deployables:** one PostgreSQL role per deployable; the AI role is read-only on the
  transactional schemas (ADR-0020, ADR-0055).
- **Secrets:** Key Vault only; nothing in images or config maps.

## Availability, backup and recovery

| Store | High availability | Backup |
|---|---|---|
| PostgreSQL | Zone-redundant standby, automatic failover | Point-in-time restore 35 days; a nightly dump per tenant to Blob (UAE) |
| Qdrant | Three nodes across zones, replication factor 2 | Snapshot per collection to Blob (UAE) |
| Broker | Three nodes across zones | The outbox is the record; nothing is lost if the broker is rebuilt |
| Redis | Premium, zone-redundant | Not backed up: it only holds what can be rebuilt |
| AKS | Nodes across zones, at least 2 replicas per deployable | Rebuilt from the registry and Terraform |

The venue edge node keeps gates and POS running through any cloud outage and syncs afterwards.

## Environments

| Environment | Purpose | Shape |
|---|---|---|
| Production | Live tenants | This document, with high availability |
| Pre-production | Acceptance tests and client demos | One small cluster, single-node stores (see the cost workbook) |
| Development | Day-to-day work | docker-compose locally (PostgreSQL 16, Redis, RabbitMQ, Qdrant) |

## Open points

- The broker is the client's choice (RabbitMQ or Kafka); if a managed service is chosen it replaces the data
  pool's broker nodes.
- Whether Qdrant runs as Qdrant Hybrid Cloud on our own AKS or plain self-hosted (ADR-0049 action).
- Prices in the workbook are estimates; confirm them in the Azure pricing calculator for UAE North.
"""


def readme():
    return f"""# HLD and LLD

> **Date:** {DATE}. Generated by `tools/build-hld-lld.py`; re-run it after any change to the deployables,
> the Terraform cell module or the decisions it quotes.

| File | What it is |
|---|---|
| [`TICVAI-HLD.md`](TICVAI-HLD.md) and [`TICVAI-HLD.svg`](TICVAI-HLD.svg) | High-level design: the cloud, the venue site and the outside world, every connection typed |
| [`TICVAI-LLD.md`](TICVAI-LLD.md) and [`TICVAI-LLD.svg`](TICVAI-LLD.svg) | Low-level design on Azure: region, VNet and subnets, tiers, security, availability, environments |
| `TICVAI - Azure Cloud Specs & Cost.xlsx` | Specs and monthly cost, without and with high availability, for production and pre-production; editable rate, margin and prices |

The polished drawings are made in Claude Design from [`../design-batches/HLD-LLD/`](../design-batches/HLD-LLD/BRIEF.md).
"""


def design_brief():
    return f"""# Claude Design: TICVAI HLD and LLD

> **Date:** {DATE}. The data in `architecture.json` is generated by `tools/build-hld-lld.py`; do not edit it
> by hand.

## What to make

**One working file: `return/TICVAI Architecture.dc.html`**, with two pages:

1. **`#hld`: the high-level design.** The classic style: grey 3D boxes, a pale yellow zone for the Azure cloud,
   a pale blue zone for the venue site with orange sub-zones for onsite sales and access control, the
   outside world on the right, and a CONNECTIONS legend. Every line is coloured and dashed by its type
   (`links[].type` in `architecture.json`, legend in `linkTypes`). The draft to improve is
   `../../hld-lld/TICVAI-HLD.svg`.
2. **`#lld`: the Azure low-level design.** The Azure style: the region as a green outline, the VNet inside it,
   tiers left to right (edge, ingress, AKS pools, data tier), Azure service icons, NSG shields on each inbound
   path, public IPs only where `architecture.json` says so, and a side panel for operations (monitor, backup,
   Defender). Venue site on the right with the Internet cloud between. The draft is `../../hld-lld/TICVAI-LLD.svg`.

## Reference

Chinmay supplied two samples of the format on 30 September (a classic HLD and an Azure LLD for another
operator). **They are attached to the session, not stored here.** Match their visual language; take no names,
logos or content from them. Use TICVAI's names only, and no client names.

## Rules

- Every box and line comes from `architecture.json`. Do not add a component, and do not drop one; if something
  looks missing, list it in `return/FINDINGS.md`.
- Each page fits a 16:9 slide at 1920 x 1080 and prints on A3 landscape.
- The legend lists only the connection types the page uses.
- No external requests except Google Fonts. Azure icons are drawn inline as simple SVG, not fetched.
- Also export each page as `return/TICVAI-HLD.svg` and `return/TICVAI-LLD.svg`.

## When it comes back

Copy the two SVGs over the drafts in `handoff/hld-lld/` only after Chinmay has approved them, and keep
`tools/build-hld-lld.py` as the source of the data.
"""


def architecture_json():
    return json.dumps({
        "date": DATE,
        "linkTypes": {k: {"label": v[0], "color": v[1], "dash": v[2]} for k, v in LINK.items()},
        "hld": {
            "zones": [dict(zip(["label", "x", "y", "w", "h", "fill", "stroke"], z)) for z in HLD_ZONES],
            "nodes": [dict(id=k, label=v[0], sub=v[1], x=v[2], y=v[3], w=v[4], h=v[5]) for k, v in HLD_NODES.items()],
            "links": [dict(source=a, target=b, type=t, via=v) for a, b, t, v in HLD_LINKS],
        },
        "lld": {
            "region": "UAE North, zones 1-3, VNet 10.20.0.0/16",
            "nodes": [dict(id=k, label=v[0], lines=v[1], x=v[2], y=v[3], w=v[4], h=v[5]) for k, v in LLD_NODES.items()],
            "frames": {k: dict(label=v[0], x=v[1], y=v[2], w=v[3], h=v[4]) for k, v in LLD_FRAMES.items()},
            "links": [dict(source=a, target=b, type=t, via=v) for a, b, t, v in LLD_LINKS],
            "subnets": [dict(zip(["name", "cidr", "holds", "rule"], s)) for s in SUBNETS],
        },
        "deployables": {k: {"what": v, "modules": sorted(MODULES.get(k, []))} for k, v in DEPLOYABLES.items()},
    }, indent=1, ensure_ascii=False) + "\n"


def main():
    write(OUT / "TICVAI-HLD.svg", hld_svg())
    write(OUT / "TICVAI-LLD.svg", lld_svg())
    write(OUT / "TICVAI-HLD.md", hld_md())
    write(OUT / "TICVAI-LLD.md", lld_md())
    write(OUT / "README.md", readme())
    build_workbook(OUT / "TICVAI - Azure Cloud Specs & Cost.xlsx")
    write(DESIGN / "BRIEF.md", design_brief())
    write(DESIGN / "architecture.json", architecture_json())
    (DESIGN / "return").mkdir(parents=True, exist_ok=True)
    keep = DESIGN / "return" / ".gitkeep"
    if not keep.exists():
        write(keep, "")
    print("hld-lld: 2 diagrams, 2 documents, 1 workbook; Claude Design brief and architecture.json")
    for k, v in CROSS.items():
        for c in v:
            print(f"ERROR {k}: {c}")
    if any(CROSS.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
