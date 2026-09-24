import openpyxl, json, io, sys, collections
p = r"D:\Chinmay\adam\ticvai\sources\requirements\Ticvai_matrix_20260621_2.xlsx"
out = sys.argv[1]
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("sheets:", [repr(w.title) for w in wb.worksheets])
ws = wb.worksheets[0]
rows = []
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
    vals = [("" if c is None else str(c).strip()) for c in list(r)[:7]]
    vals += [""] * (7 - len(vals))
    d = dict(zip(["domain_id","domain","subdomain_id","subdomain","req_id","requirement","details"], vals))
    d["row"] = i
    if d["req_id"] or d["requirement"]:
        rows.append(d)
with io.open(out, "w", encoding="utf-8") as fh:
    for d in rows:
        fh.write(json.dumps(d, ensure_ascii=False) + "\n")
print("rows:", len(rows), "->", out)
dom = collections.Counter((d["domain_id"], d["domain"]) for d in rows)
for (i, n), c in sorted(dom.items(), key=lambda x: (int(x[0][0]) if x[0][0].isdigit() else 999)):
    print(f"  {i:>4} {n[:48]:50} {c}")
# other sheets
for w in wb.worksheets[1:]:
    print("OTHER SHEET:", w.title, w.max_row - 1)
