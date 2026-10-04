# -*- coding: utf-8 -*-
"""The complete build plan in two-week sprints and four blocks (the replan of 1 October 2026), as data, workbooks
and the presentation source.

Reads what the plan generators already decided -- every task of every block with its app-module, block, owner and
build order (handoff/service-docs/plan-tasks.csv, tools/build-service-docs.py), and the schedule of every task
(handoff/service-docs/block-a-schedule.json, tools/derive-block-a-schedule.py) -- plus the screens and contracts
(for the business module of each task), the trace (requirements per module) and the flows (which block can test
which journey end to end). Writes:

  handoff/build-plan.json                    the plan as data (for the page and the deck)
  handoff/TICVAI - Build Plan.xlsx           the workbook: blocks, sprints, packages, modules, Gantt, people
  handoff/TICVAI - Sprint Plan.xlsx          sprints, tasks by sprint, blocks, people, app-modules, flows
  docs/active/build-plan-presentation.md     slide-by-slide source for the presentation

**Nothing here is estimated by hand, and nothing is scheduled here.** Since 1 October every block is planned task
by task by the generators (sizes from Block A's formulas, the scheduler in tools/sprint_plan.py), so this file only
adds up: the deck, the tickets and the schedule cannot disagree. The pace is Block A's planned pace, replaced by
the measured pace after the first sprints (re-run, not re-estimate).
"""
import collections
import csv
import datetime as dt
import glob
import io
import json
import math
import os
import re
import sys

import yaml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sprint_plan as sp  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.environ.get("PLAN_OUT") or "."
os.chdir(ROOT)

START, PLAN_END, HOLIDAYS = sp.START, sp.PLAN_END, sp.HOLIDAYS
HOURS_PER_DAY = sp.HOURS_PER_DAY
PLAN_PACE = sp.PLAN_PACE
SPRINTS = sp.SPRINTS
MODULES, PACKAGES, PACKAGE_OF = sp.MODULES, sp.PACKAGES, sp.PACKAGE_OF
FOUNDATION = sp.FOUNDATION
NOT_COUNTED = "Chinmay Parab (lead) is not counted in capacity."
# The plan this one replaces (30 September: nine three-week sprints, Blocks A and B1-B3).
PREVIOUS = {"date": "30 September 2026", "finish": "2027-04-26", "overtimeHours": 1293,
            "blockA": "screens by 20 November, back end wired by 25 December"}
AI_REVIEW = "docs/active/ai-functions-review-30-september.json"
AI_PEOPLE = ("Kalpita Mejari", "Second AI engineer")
OPS_RE = re.compile(r"Operations(?: beyond the first release)?: ([^.]+)\.")
TABLES_RE = re.compile(r"Tables: (.+?)\. Source")
SCREEN_RE = re.compile(r"(?:APP-[A-Z]+|VM)-([A-Z]+-\d{3,}[A-Z]?)(?:-REST)?$")


def load_yaml(p):
    return yaml.safe_load(io.open(p, encoding="utf-8")) or {}


def _iso(d):
    return d.isoformat() if isinstance(d, dt.date) else d


def main():
    team = json.load(io.open("docs/active/team.json", encoding="utf-8"))
    people, _, _ = sp.load_team(team)
    pace_of = {p["name"]: p["pace"] for p in people}
    settings = sp.sprint_settings(team)
    tests = sp.test_settings(team)
    rows = list(csv.DictReader(io.open("handoff/service-docs/plan-tasks.csv", encoding="utf-8")))
    sched = json.load(io.open("handoff/service-docs/block-a-schedule.json", encoding="utf-8"))
    extra = json.load(io.open("docs/active/block-a-extra-tasks.json", encoding="utf-8"))
    days_ex = {t["key"]: float(t["days"]) for t in extra["tasks"] if t.get("days")}
    hpp = HOURS_PER_DAY / PLAN_PACE
    by_key = {r["key"]: r for r in rows}
    tasks = [r for r in rows if r["type"] == "Task"]
    feats = [r for r in rows if r["type"] == "Feature"]
    kids = collections.defaultdict(list)
    for r in tasks:
        kids[r["parent"]].append(r)

    # ---------------------------------------------------------------- what each task builds, and its module
    op_contract = {}
    for f in glob.glob("contracts/*/*.yaml"):
        c = os.path.basename(f)[:-5]
        for it in (load_yaml(f).get("paths") or {}).values():
            for o in (it or {}).values():
                if isinstance(o, dict) and o.get("operationId"):
                    op_contract[o["operationId"]] = c
    op_module = {o: sp.MODULE_OF_CONTRACT.get(c, "Platform Operations") for o, c in op_contract.items()}
    screens, platforms = {}, {}
    for f in sorted(glob.glob("screens/P*.yaml")):
        d = load_yaml(f)
        code = d["platform"]["code"]
        platforms[code] = sp.PLATFORM_NAME.get(code, code)
        for s in d["screens"]:
            ops = [a["operationId"] for a in s.get("apis") or [] if isinstance(a, dict) and a.get("operationId")]
            mods = collections.Counter(op_module[o] for o in ops if o in op_module)
            screens[s["id"]] = {"platform": code, "module": mods.most_common(1)[0][0] if mods
                                else sp.PLATFORM_MODULE.get(code, FOUNDATION)}

    def builds(r):
        """(screens, operations, tables) a task builds."""
        scr, opl, tbl = set(), set(), set()
        if r["track"] == "Frontend":
            m = SCREEN_RE.search(r["key"])
            if m:
                scr.add(m.group(1))
        m = OPS_RE.search(r["description"] or "")
        if r["track"] == "Backend" and m:
            opl |= {o.strip() for o in m.group(1).split(",") if o.strip()}
        m = TABLES_RE.search(r["description"] or "")
        if r["track"] == "Database" and m:
            tbl |= {t.strip() for t in m.group(1).split(",") if t.strip()}
        return scr, opl, tbl

    def module_of(r):
        scr, opl, tbl = builds(r)
        if scr:
            return screens.get(next(iter(scr)), {}).get("module", FOUNDATION)
        if opl:
            c = collections.Counter(op_module.get(o) for o in opl if o in op_module)
            if c:
                return c.most_common(1)[0][0]
        if r["track"] == "Database":
            m = re.match(r"(?:VM-)?MIG-([A-Z0-9_]+)", r["key"])
            schema = m.group(1).lower() if m else ""
            c = sp.SCHEMA_CONTRACT.get(schema, schema)
            return sp.MODULE_OF_CONTRACT.get(c or "", FOUNDATION)
        if r["track"] == "AI" or r["key"].startswith("AI-"):
            return "AI & Intelligence"
        return FOUNDATION

    # ---------------------------------------------------------------- hours and the calendar of every task
    a_ops = set(json.load(io.open("handoff/delivery-slice.json", encoding="utf-8"))["operations"])
    a_backend_pts = sum(int(r["points"] or 0) for r in tasks if r["block"] == "A" and r["track"] == "Backend"
                        and r["key"].startswith("SVC-"))
    per_op = a_backend_pts / max(len(a_ops), 1)
    items = []
    for r in tasks:
        k = r["key"]
        who = sched["assign"].get(k)
        days = float(r.get("days") or 0) or days_ex.get(k) or 0
        pts = float(r["points"] or 0)
        if days:
            hours, dur = days * HOURS_PER_DAY, days
        else:
            # the scheduled working days (they follow the pace model, team.json sprintPlan.pace), in hours of effort
            dur = float((sched.get("duration") or {}).get(k) or 0) or (
                pts / (PLAN_PACE * pace_of.get(who, 1.0)) if pts else 0.25)
            hours = dur * HOURS_PER_DAY * pace_of.get(who, 1.0) if pts else 0.25 * HOURS_PER_DAY
        s0 = float(sched["start"].get(k, 0))
        kind = ("block test" if k.startswith("TEST-BLOCK-") else "module test" if r["track"] == "Test"
                else "AI engine" if r["track"] == "AI" else "build")
        mod = module_of(r) if kind == "build" or kind == "AI engine" else None
        items.append({"key": k, "who": who, "block": r["block"], "am": r["parent"], "track": r["track"], "kind": kind,
                      "points": pts, "days": days, "hours": hours, "s": s0, "e": s0 + dur,
                      "end": float(sched["end"].get(k, s0 + dur)), "module": mod, "tier": int(r["tier"] or 0),
                      "seq": int(r["sequence"] or 0), "ticketed": r.get("ticketed") == "yes"})
    for it in items:                                      # a test belongs to the module of what it tests
        if it["module"] is None:
            ch = [x for x in items if x["am"] == it["am"] and x["kind"] in ("build", "AI engine")]
            c = collections.Counter()
            for x in ch:
                c[x["module"]] += x["hours"]
            it["module"] = c.most_common(1)[0][0] if c else FOUNDATION
    end_idx = sp.DAY_INDEX[max(d for d in sp.DAYS if d <= PLAN_END)] + 1

    def spread(it):
        """Hours of a task per working day of its busy time: [(day index, hours)]."""
        s_, e_ = it["s"], it["e"]
        span = max(e_ - s_, 1e-6)
        out, i = [], s_
        while i < e_ - 1e-9:
            j = min(math.floor(i) + 1, e_)
            out.append((int(math.floor(i)), it["hours"] * (j - i) / span))
            i = j
        return out or [(int(s_), it["hours"])]

    load = collections.defaultdict(collections.Counter)            # person -> sprint -> hours
    after_end = collections.Counter()
    last_day = collections.defaultdict(float)
    for it in items:
        if not it["who"]:
            continue
        last_day[it["who"]] = max(last_day[it["who"]], it["e"])
        for i, h in spread(it):
            load[it["who"]][sp.sprint_of_index(i)] += h
            if i >= end_idx:
                after_end[it["who"]] += h
    n_last = max([SPRINTS[-1]["n"]] + [n for w in load for n in load[w]])
    cal = [s for s in sp.SPRINTS_ALL if s["n"] <= n_last]

    # ---------------------------------------------------------------- blocks, app-modules and flows
    sb = sched.get("blocks") or {}
    screen_block, op_block = {}, {}
    rank = {b: i for i, b in enumerate(sp.BLOCKS)}
    for r in tasks:
        scr, opl, _ = builds(r)
        for x in scr:
            if x not in screen_block or rank[r["block"]] > rank[screen_block[x]]:
                screen_block[x] = r["block"]      # a setup screen is complete with its rest
        for o in opl:
            op_block.setdefault(o, r["block"])
    flows = [load_yaml(f) for f in sorted(glob.glob("flows/F*.yaml"))]
    claims = sp.flow_claims(flows, screen_block, op_block)

    ams = []
    am_of = {}
    for f in feats:
        ch = [x for x in items if x["am"] == f["key"]]
        build_ = [x for x in ch if x["kind"] in ("build", "AI engine")]
        scr, opl, tbl = set(), set(), set()
        for x in kids[f["key"]]:
            a, b, c = builds(x)
            scr |= a
            opl |= b
            tbl |= c
        mods = collections.Counter()
        for x in build_:
            mods[x["module"]] += x["hours"]
        mod = mods.most_common(1)[0][0] if mods else FOUNDATION
        test = next((x for x in ch if x["kind"] == "module test"), None)
        s_ = min((x["s"] for x in ch), default=0.0)
        e_ = max((x["end"] for x in ch), default=0.0)
        bs = [x["e"] for x in ch if x["track"] in ("Backend", "Database")]
        a = {"key": f["key"], "name": f["subject"], "block": f["block"], "platform": f["platform"] or "",
             "module": mod, "package": PACKAGE_OF.get(mod, "Platform Foundation"),
             "screens": len(scr), "ops": len(opl), "tables": len(tbl),
             "tasks": len([x for x in ch if x["kind"] != "module test"]),
             "ticketedTasks": len([x for x in ch if x["ticketed"] and x["kind"] != "module test"]),
             "points": round(sum(x["points"] for x in build_)), "hours": round(sum(x["hours"] for x in build_)),
             "testHours": round(test["hours"]) if test else 0, "tester": test["who"] if test else "",
             "testKey": test["key"] if test else "", "lead": f["assignee"],
             "start": sp.day(s_), "done": sp.day(max(e_ - 1e-6, 0)),
             "sprintFrom": sp.sprint_of_index(s_), "sprintTo": sp.sprint_of_index(max(e_ - 1e-6, 0)),
             "backEndFrom": sp.sprint_of_index(min(bs)) if bs else None,
             "order": int(f["sequence"] or 0)}
        ams.append(a)
        am_of[f["key"]] = a
    ams.sort(key=lambda a: (rank[a["block"]], a["order"]))

    def first_sprint(its):
        """The sprint a block's work really starts in: where a tenth of its hours have started (front ends run ahead
        against the mock server, so its first task is often sprints earlier)."""
        its = sorted((i for i in its if i["who"]), key=lambda i: i["s"])
        total, run = sum(i["hours"] for i in its) or 1.0, 0.0
        for i in its:
            run += i["hours"]
            if run >= 0.1 * total:
                return sp.sprint_of_index(i["s"])
        return 1

    # **Where Block A can end** (asked on 1 October: Sprint 3 or Sprint 4?): for each candidate sprint, the hours of
    # Block A's work (its tickets and module tests) still running when that sprint's block-test window opens, per
    # person -- the overtime it takes to finish before the window.
    a_options = []
    for n in (3, 4, 5, 6):
        cut = sp.window_of(n, tests["days"])[0]
        per = collections.Counter()
        for it in items:
            if it["block"] != "A" or it["kind"] == "block test" or not it["who"] or it["e"] <= cut:
                continue
            per[it["who"]] += it["hours"] * (it["e"] - max(it["s"], cut)) / max(it["e"] - it["s"], 1e-6)
        a_options.append({"sprint": n, "endsOn": sp.SPRINTS_ALL[n - 1]["end"].isoformat(),
                          "workDoneBy": sp.day(cut - 1).isoformat(), "overtimeHours": round(sum(per.values())),
                          "byPerson": {w: round(h) for w, h in per.most_common() if h >= 1}})

    def target_ot(b):
        """Hours of a block's work (its tickets and module tests; the AI engine of later blocks apart) still running
        when its target sprint's block-test window opens, per person: the overtime it takes to end on target."""
        cut = sp.window_of(settings["targets"][b], tests["days"])[0]
        per = collections.Counter()
        for it in items:
            if it["block"] != b or it["kind"] == "block test" or not it["who"] or it["e"] <= cut:
                continue
            if it["kind"] == "AI engine" and b != "A":
                continue
            per[it["who"]] += it["hours"] * (it["e"] - max(it["s"], cut)) / max(it["e"] - it["s"], 1e-6)
        return per

    blocks = []
    for b in sp.BLOCKS:
        mine = [a for a in ams if a["block"] == b]
        if not mine:
            continue
        x = sb.get(b, {})
        its = [i for i in items if i["block"] == b]
        cl = [fid for fid, c in claims.items() if c["complete"] == b]
        pt = [fid for fid, c in claims.items() if c["partlyFrom"] == b]
        plats = collections.Counter(a["platform"] or "platform" for a in mine)
        n = x.get("endSprint") or settings["targets"][b]
        pair = sp.test_pair(b, tests)
        blocks.append({
            "block": b, "targetSprint": settings["targets"][b], "endSprint": n,
            "endsOn": x.get("endsOn"), "targetEndsOn": x.get("targetEndsOn"), "lastWork": x.get("lastDay"),
            "testFrom": x.get("testFrom"), "testTo": x.get("testTo"), "testers": pair, "testLead": tests["lead"],
            "aiEngineLastDay": x.get("aiEngineLastDay"),
            "normalEndSprint": x.get("normalEndSprint") or n, "normalEndsOn": x.get("normalEndsOn") or x.get("endsOn"),
            "fixed": bool(x.get("fixed")),
            "targetOvertimeHours": round(sum(target_ot(b).values())),
            "targetOvertimeByPerson": {w: round(h) for w, h in target_ot(b).most_common() if h >= 1},
            "firstSprint": first_sprint(its),
            "appModules": [a["name"] for a in mine], "appModuleCount": len(mine),
            "platforms": dict(plats.most_common()),
            "screens": sum(a["screens"] for a in mine), "ops": sum(a["ops"] for a in mine),
            "tables": sum(a["tables"] for a in mine),
            "hours": round(sum(i["hours"] for i in its if i["kind"] == "build")),
            "aiHours": round(sum(i["hours"] for i in its if i["kind"] == "AI engine")),
            "testHours": round(sum(i["hours"] for i in its if i["kind"] in ("module test", "block test"))),
            "tasks": len([i for i in its if i["kind"] != "block test"]),
            "ticketedTasks": len([i for i in its if i["ticketed"]]),
            "flowsClaimed": len(cl), "flowsPartly": len(pt),
            "criticalFlows": sorted(fid for fid in cl if claims[fid]["criticality"] in ("revenue", "safety", "financial")),
            "testScope": (f"End to end on {', '.join(p for p in plats if p != 'platform') or 'the platform'}: "
                          f"{sum(a['screens'] for a in mine)} screens, {sum(a['ops'] for a in mine)} operations, "
                          f"{sum(a['tables'] for a in mine)} tables in {len(mine)} app-modules; {len(cl)} flows "
                          f"end to end, {len(pt)} more partly; offline runs for POS, scanner and staff app where "
                          "present; a load run where the block adds a purchase or admission path.")})

    # ---------------------------------------------------------------- modules and packages
    mods = collections.OrderedDict((m, {"module": m, "package": PACKAGE_OF[m], "requirements": 0,
                                        "requirementsWaiting": 0, "hoursByBlock": collections.Counter(),
                                        "screensByBlock": collections.Counter(), "opsByBlock": collections.Counter(),
                                        "people": collections.Counter(), "s": None, "e": None}) for m in MODULES)
    trace = json.load(io.open("handoff/traceability.json", encoding="utf-8"))["rows"]
    for r in trace:
        m = sp.MODULE_OF_CONTRACT.get(r.get("contract") or "")
        if m and r.get("verdict") in ("CONTRACTED", "CONTRACTED_PARTIAL"):
            mods[m]["requirements"] += 1
        if m and r.get("verdict") == "CONTRACTED_PARTIAL":
            mods[m]["requirementsWaiting"] += 1
    for it in items:
        m = mods.get(it["module"]) or mods[FOUNDATION]
        m["hoursByBlock"][it["block"]] += it["hours"]
        if it["who"]:
            m["people"][it["who"]] += it["hours"]
        m["s"] = it["s"] if m["s"] is None else min(m["s"], it["s"])
        m["e"] = it["end"] if m["e"] is None else max(m["e"], it["end"])
    for r in tasks:
        scr, opl, _ = builds(r)
        for x in scr:
            mods.get(screens.get(x, {}).get("module"), mods[FOUNDATION])["screensByBlock"][r["block"]] += 1
        for o in opl:
            mods.get(op_module.get(o), mods[FOUNDATION])["opsByBlock"][r["block"]] += 1
    mod_rows = []
    for m in mods.values():
        top = m["people"].most_common()
        mod_rows.append({"module": m["module"], "package": m["package"], "requirements": m["requirements"],
                         "requirementsWaiting": m["requirementsWaiting"],
                         "hoursByBlock": {b: round(m["hoursByBlock"][b]) for b in sp.BLOCKS},
                         "screensByBlock": {b: m["screensByBlock"][b] for b in sp.BLOCKS},
                         "opsByBlock": {b: m["opsByBlock"][b] for b in sp.BLOCKS},
                         "hours": round(sum(m["hoursByBlock"].values())),
                         "screens": sum(m["screensByBlock"].values()), "ops": sum(m["opsByBlock"].values()),
                         "start": sp.day(m["s"]) if m["s"] is not None else None,
                         "end": sp.day(max(m["e"] - 1e-6, 0)) if m["e"] is not None else None,
                         "sprints": (f"{sp.sprint_of_index(m['s'])}-{sp.sprint_of_index(max(m['e'] - 1e-6, 0))}"
                                     if m["s"] is not None else ""),
                         "appModules": len([a for a in ams if a["module"] == m["module"]]),
                         "lead": top[0][0] if top else "", "team": [n for n, _ in top[:6]]})
    pk = []
    for name, why, ms in PACKAGES:
        rs = [x for x in mod_rows if x["module"] in ms]
        ppl = collections.Counter()
        for x in rs:
            ppl.update({n: 1 for n in x["team"]})
        starts = [x["start"] for x in rs if x["start"]]
        ends = [x["end"] for x in rs if x["end"]]
        pk.append({"package": name, "why": why, "modules": ms, "hours": sum(x["hours"] for x in rs),
                   "hoursByBlock": {b: sum(x["hoursByBlock"][b] for x in rs) for b in sp.BLOCKS},
                   "requirements": sum(x["requirements"] for x in rs), "screens": sum(x["screens"] for x in rs),
                   "ops": sum(x["ops"] for x in rs), "start": min(starts) if starts else None,
                   "end": max(ends) if ends else None,
                   "lead": max(rs, key=lambda x: x["hours"])["lead"] if rs else "",
                   "appModules": sum(x["appModules"] for x in rs)})

    # ---------------------------------------------------------------- people and sprints
    people_rows = []
    for p in people:
        name, frm = p["name"], p["from"]
        cap = {sp_["n"]: len(sp.workdays(max(sp_["start"], frm), sp_["end"])) * HOURS_PER_DAY * p["pace"] for sp_ in cal}
        weeks = max(1, len([d for d in sp.DAYS if frm <= d <= PLAN_END]) / 5)
        a_pts = sum(it["points"] for it in items if it["who"] == name and it["block"] == "A" and it["kind"] == "build")
        a_end = max([it["e"] for it in items if it["who"] == name and it["block"] == "A"] or [0.0])
        people_rows.append({"name": name, "role": p["role"], "pace": p["pace"], "from": frm,
                            "capacity": {n: round(v) for n, v in cap.items()},
                            "planned": {n: round(load[name][n]) for n in sorted(load[name])},
                            "hours": round(sum(load[name].values())),
                            "testHours": round(sum(it["hours"] for it in items if it["who"] == name
                                                   and it["kind"] in ("module test", "block test"))),
                            "lastDay": sp.day(max(last_day[name] - 1e-6, 0)) if last_day.get(name) else None,
                            "hoursAfterPlanEnd": round(after_end[name]),
                            "overtimeHoursPerWeek": round(after_end[name] / weeks, 1),
                            "blockAPoints": round(a_pts), "blockAEnd": sp.day(max(a_end - 1e-6, 0)) if a_end else None,
                            "modules": [m for m, _ in collections.Counter(
                                it["module"] for it in items if it["who"] == name).most_common(8)]})
    sprints = []
    for s in cal:
        its = [it for it in items if it["who"] and sp.sprint_of_index(it["s"]) == s["n"]]
        ending = [b["block"] for b in blocks if b["endSprint"] == s["n"]]
        active = sorted({it["block"] for it in its}, key=lambda b: rank[b])
        mt = [it for it in its if it["kind"] == "module test"]
        hol = [f"{_d(d)} {n}" for d, n in HOLIDAYS.items() if s["start"] <= d <= s["end"]]
        sprints.append({"n": s["n"], "name": s["name"], "start": s["start"], "end": s["end"], "days": s["days"],
                        "inPlan": s["inPlan"],
                        "capacity": round(sum(r["capacity"].get(s["n"], 0) for r in people_rows)),
                        "planned": round(sum(load[w][s["n"]] for w in load)),
                        # developers only (the AI engineers apart): capacity left unplanned in the sprint
                        "devIdle": round(sum(r["capacity"].get(s["n"], 0) - load[r["name"]][s["n"]]
                                             for r in people_rows if r["name"] not in AI_PEOPLE)),
                        "blocks": active, "blockEnding": ending,
                        "blockTest": [{"block": b["block"], "from": b["testFrom"], "to": b["testTo"]}
                                      for b in blocks if b["endSprint"] == s["n"]],
                        "moduleTests": len(mt), "moduleTestHours": round(sum(it["hours"] for it in mt)),
                        "appModulesDone": [a["name"] for a in ams if a["sprintTo"] == s["n"]],
                        "holidays": hol, "buffer": s["n"] in settings["bufferSprints"]})

    # ---------------------------------------------------------------- build phases, back end and front end
    BACK = ("Backend", "Database", "DevOps", "Setup", "AI")
    ph = collections.OrderedDict()
    for it in items:
        if it["kind"] not in ("build", "AI engine") or it["key"].startswith("SETUP-ONBOARD"):
            continue
        side = "back end" if it["track"] in BACK else "front end"
        k = (it["tier"], side)
        r = ph.setdefault(k, {"phase": it["tier"], "name": sp.PHASE_NAME.get(it["tier"], ""), "side": side,
                              "hours": 0.0, "byBlock": collections.Counter(), "s": it["s"], "e": it["end"],
                              "modules": collections.Counter(), "people": collections.Counter()})
        r["hours"] += it["hours"]
        r["byBlock"][it["block"]] += it["hours"]
        r["s"], r["e"] = min(r["s"], it["s"]), max(r["e"], it["end"])
        r["modules"][it["module"]] += it["hours"]
        if it["who"]:
            r["people"][it["who"]] += it["hours"]
    phases = [{"phase": r["phase"], "name": r["name"], "side": r["side"], "hours": round(r["hours"]),
               "hoursByBlock": {b: round(r["byBlock"][b]) for b in sp.BLOCKS},
               "start": sp.day(r["s"]).isoformat(), "end": sp.day(max(r["e"] - 1e-6, 0)).isoformat(),
               "modules": [m for m, _ in r["modules"].most_common(8)],
               "people": [n for n, _ in r["people"].most_common(6)]} for _, r in sorted(ph.items())]

    # ---------------------------------------------------------------- totals
    dev = [it for it in items if it["kind"] == "build"]
    devs_finish = max(it["e"] for it in dev if it["who"])
    ai_finish = max([it["e"] for it in items if it["kind"] == "AI engine"] or [0.0])   # unassigned ones included
    test_items = [it for it in items if it["kind"] in ("module test", "block test")]
    counted = [p for p in people_rows if p["name"] not in AI_PEOPLE]
    # Block D against its target, and the developers' spare capacity once their planned work is done (1 October:
    # it absorbs Claude Design returns, defects, change requests, and helps the AI developers)
    xd = next((x for x in blocks if x["block"] == "D"), None)
    block_d = ({"onTarget": xd["normalEndSprint"] <= xd["targetSprint"], "overtimeHours": xd["targetOvertimeHours"],
                "normalEndsOn": xd["normalEndsOn"], "targetEndsOn": xd["targetEndsOn"]} if xd else {})
    bs = [s_ for s_ in sprints if s_.get("buffer")]
    buffer = {"sprints": [s_["n"] for s_ in bs], "from": bs[0]["start"].isoformat() if bs else None,
              "to": bs[-1]["end"].isoformat() if bs else None, "hours": round(sum(s_["devIdle"] for s_ in bs)),
              "hoursPerSprint": round(sum(s_["devIdle"] for s_ in bs) / len(bs)) if bs else 0}
    fin_n = sp.sprint_of_index(max(devs_finish - 1e-6, 0))
    after = [s_ for s_ in sprints if s_["inPlan"] and s_["n"] > fin_n]
    spare = {"from": sp.day(max(devs_finish - 1e-6, 0)).isoformat(), "fromSprint": fin_n + 1 if after else None,
             "sprints": len(after), "hours": round(sum(s_["devIdle"] for s_ in after)),
             "hoursPerSprint": round(sum(s_["devIdle"] for s_ in after) / len(after)) if after else 0}
    # **Block A's critical path** (asked 1 October): its longest chain of waits with every task on its own person
    # (no one waits for a free person), at the scheduled durations; soft waits hold only the finish. The shortest
    # Block A can take whatever the team size.
    a_its = sorted((it for it in items if it["block"] == "A" and it["kind"] != "block test"), key=lambda it: it["seq"])
    a_keys_ = {it["key"] for it in a_its}
    ef = {}
    for it in a_its:
        r = by_key[it["key"]]
        st = 0.0
        wire = 0.0
        for d in (r["dependsOn"] or "").split():
            if d not in a_keys_ or d not in ef:
                continue
            soft = sp.is_soft({"track": r["track"], "key": r["key"], "service": r["service"]},
                              {"track": by_key[d]["track"], "key": d, "service": by_key[d]["service"]})
            if soft:
                wire = max(wire, ef[d])
            else:
                st = max(st, ef[d])
        ef[it["key"]] = max(st + (it["e"] - it["s"]), wire)
    a_critical = round(max(ef.values() or [0.0]), 1)
    out = {
        "generatedBy": "tools/build-plan-deck.py", "generated": dt.date.today().isoformat(),
        "basis": {"pacePointsPerDeveloperDay": round(PLAN_PACE, 2), "hoursPerPoint": round(hpp, 3),
                  "pointsPerOperation": round(per_op, 2),
                  "totalPoints": round(sum(it["points"] for it in items)),
                  "totalHours": round(sum(it["hours"] for it in items)),
                  "buildHours": round(sum(it["hours"] for it in dev)),
                  "testHours": round(sum(it["hours"] for it in test_items)),
                  "moduleTestHours": round(sum(it["hours"] for it in test_items if it["kind"] == "module test")),
                  "blockTestHours": round(sum(it["hours"] for it in test_items if it["kind"] == "block test")),
                  "aiEngineHours": round(sum(it["hours"] for it in items if it["kind"] == "AI engine")),
                  "forecastFinish": sp.day(max(devs_finish - 1e-6, 0)).isoformat(),
                  "aiFinish": sp.day(max(ai_finish - 1e-6, 0)).isoformat(),
                  "planEnd": PLAN_END.isoformat(),
                  "overtimeHours": round(sum(p["hoursAfterPlanEnd"] for p in people_rows)),
                  "overtimeHoursDevelopers": round(sum(p["hoursAfterPlanEnd"] for p in counted)),
                  "notCounted": NOT_COUNTED, "previous": PREVIOUS,
                  "tasks": len(tasks), "ticketedTasks": len([r for r in tasks if r.get("ticketed") == "yes"]),
                  "appModules": len(ams), "blockAOptions": a_options,
                  "pace": sched.get("pace") or "", "blockACriticalPathDays": a_critical,
                  "aiUnassignedHours": round(sum(it["hours"] for it in items if it["kind"] == "AI engine"
                                                 and not it["who"])),
                  "devIdleHoursToPlanEnd": round(sum(s_["devIdle"] for s_ in sprints if s_["inPlan"])),
                  "blockD": block_d, "spare": spare, "buffer": buffer},
        "calendar": {"start": START.isoformat(), "planEnd": PLAN_END.isoformat(), "sprintDays": sp.SPRINT_DAYS,
                     "holidays": {d.isoformat(): n for d, n in HOLIDAYS.items()}},
        "blocks": blocks, "sprints": sprints, "appModules": ams, "packages": pk, "modules": mod_rows,
        "people": people_rows, "phases": phases,
        "flows": [dict(id=k, **v) for k, v in claims.items()],
        "platforms": platforms,
    }
    task_rows = []
    for it in sorted(items, key=lambda x: (sp.sprint_of_index(x["s"]) if x["who"] else 99, x["seq"])):
        r = by_key[it["key"]]
        task_rows.append({"sprint": sp.sprint_of_index(it["s"]) if it["who"] else "", "block": it["block"],
                          "appModule": by_key[r["parent"]]["subject"] if r["parent"] in by_key else "",
                          "platform": am_of[r["parent"]]["platform"] if r["parent"] in am_of else "",
                          "key": it["key"], "title": re.sub(r"^\[[^\]]+\]\s*", "", r["subject"]),
                          "track": r["track"], "kind": it["kind"], "owner": it["who"] or "",
                          "points": int(it["points"]) if it["points"] else "", "hours": round(it["hours"], 1),
                          "starts": sp.day(it["s"]).isoformat(), "ends": sp.day(max(it["end"] - 1e-6, 0)).isoformat(),
                          "depends": r["dependsOn"], "ticketed": r.get("ticketed", ""), "sequence": it["seq"]})
    return out, task_rows


def jsonable(o):
    if isinstance(o, dt.date):
        return o.isoformat()
    if isinstance(o, collections.Counter):
        return dict(o)
    if isinstance(o, set):
        return sorted(o)
    raise TypeError(type(o))


# ---------------------------------------------------------------- outputs
def _d(x):
    if not x:
        return ""
    if isinstance(x, str):
        x = dt.date.fromisoformat(x)
    return x.strftime("%d %b %Y").lstrip("0")


def _pp(n):
    return f"{int(round(n)):,}"


def _opt_a(b, a):
    """Block A's option at its target sprint (team.json sprintPlan.blocks): the overtime it takes, by person."""
    return next((o for o in b.get("blockAOptions") or [] if o["sprint"] == a["targetSprint"]),
                {"overtimeHours": 0, "byPerson": {}})


def d_text(b):
    """Block D against its target, from the plan: on target without overtime, or the hours it takes."""
    x = b.get("blockD") or {}
    if not x:
        return ""
    if x.get("onTarget"):
        return f"Block D lands on {_d(x['targetEndsOn'])} at normal hours, without overtime."
    return (f"Block D at normal hours ends {_d(x['normalEndsOn'])}; landing it on {_d(x['targetEndsOn'])} takes about "
            f"{_pp(x['overtimeHours'])} hours of overtime.")


def spare_value(b):
    x = b.get("spare") or {}
    return f"{_pp(x.get('hours', 0))} hours from {_d(x.get('from'))}" if x.get("sprints") else "none before 2 April"


def buffer_text(b):
    """Sprints 11-13 are buffer (decided 1 October): what is free in them, and what it is for."""
    x = b.get("buffer") or {}
    if not x.get("sprints"):
        return spare_text(b)
    return (f"Sprints {x['sprints'][0]}-{x['sprints'][-1]} ({_d(x['from'])} to {_d(x['to'])}) are buffer (decided 1 "
            f"October): about {_pp(x['hoursPerSprint'])} developer hours a sprint ({_pp(x['hours'])} in all) for Claude "
            "Design returns, defects, change requests and helping the AI developers.")


TICKETING = "Block A in full and the next block (B) task by task in release r2; Blocks C and D are ticketed as app-modules (features) and broken into tasks one block ahead, at each block's start, through the normal Tuesday/Friday release. Their tasks are already planned (plan-tasks.csv) with the keys they will get."


def spare_text(b):
    """The developers' capacity left once their planned work is done, to 2 April."""
    x = b.get("spare") or {}
    if not x.get("sprints"):
        return "The developers' planned work fills the six months."
    return (f"The developers' planned work is done by {_d(x['from'])}; from Sprint {x['fromSprint']} to 2 April about "
            f"{_pp(x['hoursPerSprint'])} hours a sprint ({_pp(x['hours'])} in all) are free. They absorb Claude Design "
            "returns, defects, change requests, and help the AI developers.")


def _style():
    from openpyxl.styles import Font, PatternFill
    return {"HEAD": PatternFill("solid", fgColor="1F3864"), "SUB": PatternFill("solid", fgColor="D9E1F2"),
            # A2 (Block A's second half, sprint_plan.BLOCKS) in a lighter shade of A's blue: without it the Gantt
            # stopped on the first A2 app-module (KeyError 'A2').
            "BLOCK": {"A": PatternFill("solid", fgColor="2E75B6"), "A2": PatternFill("solid", fgColor="9DC3E6"),
                      "B": PatternFill("solid", fgColor="70AD47"),
                      "C": PatternFill("solid", fgColor="FFC000"), "D": PatternFill("solid", fgColor="A5A5A5")},
            "TEST": PatternFill("solid", fgColor="C00000"), "OVER": PatternFill("solid", fgColor="F4B084"),
            "WHITE": Font(color="FFFFFF", bold=True)}


def _sheet(wb, title, header, rows, widths=None, first=False):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    st = _style()
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = st["HEAD"], st["WHITE"]
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def summary_rows(plan):
    b = plan["basis"]
    a = next(x for x in plan["blocks"] if x["block"] == "A")
    rows = [
        ["Start", _d(plan["calendar"]["start"]), "Sprint 1 starts Monday 5 October 2026."],
        ["Sprints", f"{len([s for s in plan['sprints'] if s['inPlan']])} sprints of two weeks to {_d(plan['calendar']['planEnd'])}",
         "Sprint 1 = 5-16 Oct 2026 ... Sprint 13 = 22 Mar-2 Apr 2027 (replan of 1 October; it was nine sprints of three weeks)."],
    ]
    for x in plan["blocks"]:
        rows.append([f"Block {x['block']}", f"ends {_d(x['endsOn'])} (Sprint {x['endSprint']})",
                     f"Target Sprint {x['targetSprint']} ({_d(x['targetEndsOn'])}). {x['appModuleCount']} app-modules; "
                     f"block test {_d(x['testFrom'])} to {_d(x['testTo'])}; {x['flowsClaimed']} flows end to end."
                     + (f" AI engine capabilities to {_d(x['aiEngineLastDay'])}." if x.get("aiEngineLastDay") else "")])
    rows += [
        ["Forecast finish (developers)", _d(b["forecastFinish"]),
         f"At normal hours. The plan of {PREVIOUS['date']} said {_d(PREVIOUS['finish'])}."],
        ["AI engine finish", _d(b["aiFinish"]), "Two AI engineers, no third (decided 30 September); the AI review "
         "found 50 AI-weeks against 35 available in Block B onwards. Decided 1 October: every AI engine task is "
         "created; those past 2 April are left unassigned for the AI developers joining."],
        ["Total effort", f"{_pp(b['totalHours'])} hours",
         f"Build {_pp(b['buildHours'])} h, testing {_pp(b['testHours'])} h (module tests {_pp(b['moduleTestHours'])}, "
         f"block tests {_pp(b['blockTestHours'])}), AI engine {_pp(b['aiEngineHours'])} h."],
        ["Overtime to finish by 2 April", f"{_pp(b['overtimeHoursDevelopers'])} hours (developers)",
         f"Decided 1 October: Block D keeps its scope to 2 April. {d_text(b)} The AI engine's "
         f"{_pp(b.get('aiUnassignedHours', 0))} h past 2 April is not overtime: those tasks are "
         f"unassigned for the AI developers joining. The plan of {PREVIOUS['date']} needed about "
         f"{_pp(PREVIOUS['overtimeHours'])} h without the tests."],
        ["Buffer", f"Sprints {'-'.join(str(n) for n in ((b.get('buffer') or {}).get('sprints') or [])[::max(1, len((b.get('buffer') or {}).get('sprints') or [1]) - 1)])}, "
                   f"{_pp((b.get('buffer') or {}).get('hours', 0))} hours", buffer_text(b)],
        ["Block A", f"{a['appModuleCount']} app-modules, ends {_d(a['targetEndsOn'])} (Sprint {a['targetSprint']})",
         f"Previously: {PREVIOUS['blockA']}. Now complete, tested and accepted as one block. Decided 1 October: "
         f"Sprint {a['targetSprint']} with {_pp(_opt_a(b, a)['overtimeHours'])} h of planned overtime ("
         + ", ".join(f"{n} {h}" for n, h in _opt_a(b, a)['byPerson'].items())
         + f"); at normal hours it would end {_d(a.get('normalEndsOn') or a['endsOn'])}."],
        ["Ticketing depth", f"{_pp(b['tasks'])} tasks planned, {_pp(b['ticketedTasks'])} ticketed", TICKETING],
        ["Pace", b.get("pace") or f"{b['pacePointsPerDeveloperDay']} points per developer per day",
         "Default: Block A's plan of record (3,018 points, 35 days, 9 developers), replaced by the measured pace; "
         "team.json sprintPlan.pace sets a scenario (tasks a developer-day, ramp)."],
        ["Block A's critical path", f"{b.get('blockACriticalPathDays')} working days",
         "Its longest chain of waits with every task on its own person: the shortest Block A can take at this pace, "
         "whatever the team size. Block A is 40 working days (decided 1 October; it was 35)."],
        ["Points per back-end operation", f"{b['pointsPerOperation']}", "Block A's back-end points divided by its operations."],
        ["Team", f"{len(plan['people'])} people", b["notCounted"]],
    ]
    return rows


def write_build_xlsx(plan, path):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side, Alignment
    from openpyxl.utils import get_column_letter
    st = _style()
    thin = Side(style="thin", color="BFBFBF")
    wb = Workbook()
    _sheet(wb, "Summary", ["What", "Value", "How it is worked out"], summary_rows(plan), [30, 36, 110], first=True)
    _sheet(wb, "Blocks", ["Block", "Target sprint", "Final sprint", "Ends", "Block test", "Testers", "App-modules",
                          "Screens", "Operations", "Tables", "Build hours", "Test hours", "AI engine hours",
                          "Flows end to end", "Flows partly", "Apps"],
           [[x["block"], x["targetSprint"], x["endSprint"], _d(x["endsOn"]), f"{_d(x['testFrom'])} to {_d(x['testTo'])}",
             " + ".join(x["testers"]) + f" (led by {x['testLead']})", x["appModuleCount"], x["screens"], x["ops"],
             x["tables"], x["hours"], x["testHours"], x["aiHours"], x["flowsClaimed"], x["flowsPartly"],
             ", ".join(f"{k} ({v})" for k, v in x["platforms"].items())] for x in plan["blocks"]],
           [7, 8, 8, 13, 26, 40, 9, 8, 9, 8, 9, 9, 9, 9, 9, 70])
    _sheet(wb, "Sprints", ["Sprint", "Starts", "Ends", "Working days", "Blocks in progress", "Block ending",
                           "Capacity (hours)", "Planned (hours)", "Used", "Holidays"],
           [[s["name"] + (" (buffer)" if s.get("buffer") else "") + ("" if s["inPlan"] else " (past 2 Apr)"),
             _d(s["start"]), _d(s["end"]), s["days"],
             ", ".join(s["blocks"]), ", ".join(s["blockEnding"]), s["capacity"], s["planned"],
             f"{s['planned'] / max(s['capacity'], 1):.0%}", "; ".join(s["holidays"])] for s in plan["sprints"]],
           [16, 13, 13, 8, 12, 10, 11, 11, 8, 40])
    _sheet(wb, "Packages", ["Package", "What it is", "Modules", "App-modules", "Requirements covered", "Screens",
                            "Operations", "Hours"] + [f"Block {b} hours" for b in sp.BLOCKS] + ["Starts", "Completion", "Lead"],
           [[p["package"], p["why"], len(p["modules"]), p["appModules"], p["requirements"], p["screens"], p["ops"],
             p["hours"]] + [p["hoursByBlock"][b] for b in sp.BLOCKS] + [_d(p["start"]), _d(p["end"]), p["lead"]]
            for p in plan["packages"]], [26, 60, 8, 9, 11, 8, 9, 8, 8, 8, 8, 8, 13, 13, 22])
    _sheet(wb, "Modules", ["Package", "Module", "Requirements covered", "Waiting on the client", "App-modules",
                           "Screens", "Operations"] + [f"Block {b} hours" for b in sp.BLOCKS]
           + ["Total hours", "Starts", "Completion", "Sprints", "Lead", "Team"],
           [[m["package"], m["module"], m["requirements"], m["requirementsWaiting"], m["appModules"], m["screens"],
             m["ops"]] + [m["hoursByBlock"][b] for b in sp.BLOCKS]
            + [m["hours"], _d(m["start"]), _d(m["end"]), m["sprints"], m["lead"], ", ".join(m["team"])]
            for m in plan["modules"]], [24, 32, 10, 9, 9, 8, 9, 8, 8, 8, 8, 9, 13, 13, 8, 22, 70])
    _sheet(wb, "Build phases", ["Phase", "Side", "Starts", "Ends"] + [f"Block {b} hours" for b in sp.BLOCKS]
           + ["Hours", "Main modules", "People"],
           [[f"{r['phase']} {r['name']}", r["side"], _d(r["start"]), _d(r["end"])]
            + [r["hoursByBlock"][b] for b in sp.BLOCKS] + [r["hours"], ", ".join(r["modules"]), ", ".join(r["people"])]
            for r in plan["phases"]], [16, 10, 13, 13, 8, 8, 8, 8, 8, 70, 60])

    # Gantt: one row per block then its app-modules, one column per week
    ws = wb.create_sheet("Gantt")
    weeks, d = [], START
    end = max(dt.date.fromisoformat(plan["basis"]["forecastFinish"]), PLAN_END)
    while d <= end + dt.timedelta(days=6):
        weeks.append(d)
        d += dt.timedelta(days=7)
    ws.append(["Block / app-module", "Lead", "Hours", "Done"] + [""] * len(weeks))
    ws.append(["", "", "", ""] + [w.strftime("%d %b") for w in weeks])
    for i, w in enumerate(weeks):
        sp_ = next((s for s in plan["sprints"] if s["start"] == w.isoformat()), None)
        ws.cell(row=1, column=5 + i, value=f"S{sp_['n']}" if sp_ else "")
    for c in ws[1] + ws[2]:
        c.fill, c.font = st["HEAD"], st["WHITE"]
        c.alignment = Alignment(horizontal="center", text_rotation=90 if c.row == 2 and c.column > 4 else 0)
    r = 3

    def bar(row, s0, e0, fill):
        for i, w in enumerate(weeks):
            if w <= e0 and w + dt.timedelta(days=4) >= s0:
                c = ws.cell(row=row, column=5 + i)
                c.fill = st["OVER"] if w > PLAN_END else fill
                c.border = Border(top=thin, bottom=thin)

    for x in plan["blocks"]:
        ws.cell(row=r, column=1, value=f"Block {x['block']} (Sprints {x['firstSprint']}-{x['endSprint']})").font = Font(bold=True)
        ws.cell(row=r, column=3, value=x["hours"] + x["testHours"])
        ws.cell(row=r, column=4, value=_d(x["endsOn"]))
        for c in range(1, 5 + len(weeks)):
            ws.cell(row=r, column=c).fill = st["SUB"]
        if x["testFrom"]:
            bar(r, dt.date.fromisoformat(x["testFrom"]), dt.date.fromisoformat(x["testTo"]), st["TEST"])
        r += 1
        for a in [a for a in plan["appModules"] if a["block"] == x["block"]]:
            ws.cell(row=r, column=1, value="   " + a["name"])
            ws.cell(row=r, column=2, value=a["lead"])
            ws.cell(row=r, column=3, value=a["hours"])
            ws.cell(row=r, column=4, value=_d(a["done"]))
            bar(r, dt.date.fromisoformat(_iso(a["start"])), dt.date.fromisoformat(_iso(a["done"])), st["BLOCK"][a["block"]])
            r += 1
    r += 1
    for b in sp.BLOCKS:
        ws.cell(row=r, column=1, value=f"Block {b}")
        ws.cell(row=r, column=2).fill = st["BLOCK"][b]
        r += 1
    for fill, text in ((st["TEST"], "Block test (no new feature work starts)"), (st["OVER"], "After 2 April")):
        ws.cell(row=r, column=1, value=text)
        ws.cell(row=r, column=2).fill = fill
        r += 1
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 12
    for i in range(len(weeks)):
        ws.column_dimensions[get_column_letter(5 + i)].width = 3.2
    ws.freeze_panes = "E3"
    write_people(wb, plan)
    _sheet(wb, "Release checklist", ["When", "Step", "Check", "Owner"], [list(x) for x in RELEASE_CHECKLIST], [26, 14, 90, 22])
    _sheet(wb, "Decisions", ["Decision", "Status", "What was decided", "Why", "What it costs"], [list(x) for x in ARCH_DECISIONS], [34, 10, 60, 60, 50])
    _sheet(wb, "Risks", ["Risk", "Signal", "What we do"], [list(x) for x in RISKS], [40, 50, 70])
    wb.save(path)


def write_people(wb, plan):
    st = _style()
    sprints = plan["sprints"]
    rows = []
    for pe in plan["people"]:
        rows.append([pe["name"], pe["role"], f"{pe['pace']:.0%}", _d(pe["from"]), pe["blockAPoints"], _d(pe["blockAEnd"])]
                    + [pe["planned"].get(str(s["n"]), pe["planned"].get(s["n"], 0)) for s in sprints]
                    + [pe["hours"], pe["testHours"], _d(pe["lastDay"]), pe["hoursAfterPlanEnd"], pe["overtimeHoursPerWeek"],
                       ", ".join(pe["modules"][:6])])
    ws = _sheet(wb, "People", ["Name", "Role", "Pace", "From", "Block A points", "Block A work ends"]
                + [f"S{s['n']} hours" for s in sprints]
                + ["Total hours", "Of which testing", "Last day at normal hours", "Hours past 2 Apr", "Overtime h/week to finish", "Modules"],
                rows, [26, 30, 7, 12, 8, 13] + [7] * len(sprints) + [9, 9, 13, 9, 10, 70])
    for row in ws.iter_rows(min_row=2):
        pe = next(p for p in plan["people"] if p["name"] == row[0].value)
        for i, s in enumerate(sprints):
            c = row[6 + i]
            cap = pe["capacity"].get(str(s["n"]), pe["capacity"].get(s["n"], 0))
            if isinstance(c.value, (int, float)) and c.value > cap * 1.02:
                c.fill = st["OVER"]
    return ws


def write_sprint_xlsx(plan, task_rows, path):
    """TICVAI - Sprint Plan.xlsx (asked for on 1 October): sprints, tasks by sprint, blocks, people, app-modules, flows."""
    from openpyxl import Workbook
    st = _style()
    wb = Workbook()
    _sheet(wb, "Summary", ["What", "Value", "How it is worked out"], summary_rows(plan), [30, 36, 110], first=True)
    rows = []
    for s in plan["sprints"]:
        bt = "; ".join(f"Block {x['block']} test {_d(x['from'])} to {_d(x['to'])}" for x in s["blockTest"])
        rows.append([s["name"] + (" (buffer)" if s.get("buffer") else "") + ("" if s["inPlan"] else " (past 2 Apr)"),
                     _d(s["start"]), _d(s["end"]), s["days"],
                     ", ".join(s["blocks"]), ", ".join(f"Block {b}" for b in s["blockEnding"]), bt,
                     s["capacity"], s["planned"], f"{s['planned'] / max(s['capacity'], 1):.0%}",
                     s["moduleTests"], s["moduleTestHours"], len(s["appModulesDone"]), s.get("devIdle", ""),
                     "; ".join(s["holidays"])])
    ws = _sheet(wb, "Sprints", ["Sprint", "Starts", "Ends", "Working days", "Blocks in progress", "Block ending",
                                "Block test window (no new feature work starts)", "Capacity (hours)", "Planned (hours)",
                                "Used", "Module tests", "Module test hours", "App-modules done",
                                "Developer hours unplanned", "Holidays"],
                rows, [18, 12, 12, 8, 11, 10, 34, 10, 10, 7, 8, 9, 9, 10, 36])
    for row in ws.iter_rows(min_row=2):
        if row[6].value:
            row[6].fill = st["TEST"]
            row[6].font = st["WHITE"]
    rows = [[t["sprint"], t["block"], t["appModule"], t["platform"], t["key"], t["title"], t["kind"], t["track"],
             t["owner"], t["points"], t["hours"], t["starts"], t["ends"], t["depends"], t["ticketed"]] for t in task_rows]
    ws = _sheet(wb, "Tasks by sprint", ["Sprint", "Block", "App-module", "Platform", "Task key", "Title", "Kind", "Track",
                                        "Owner", "Points", "Hours", "Starts", "Ends", "Depends on",
                                        "Ticketed in OpenProject"],
                rows, [7, 6, 40, 16, 28, 60, 11, 10, 22, 7, 7, 11, 11, 40, 9])
    for row in ws.iter_rows(min_row=2):
        if row[6].value in ("module test", "block test"):
            for c in row[:9]:
                c.fill = st["SUB"]
    rows = [[x["block"], x["firstSprint"], x["endSprint"], _d(x["endsOn"]), x["targetSprint"], _d(x["targetEndsOn"]),
             f"Sprint {x.get('normalEndSprint')}, {_d(x.get('normalEndsOn'))}", x.get("targetOvertimeHours", 0),
             ", ".join(f"{k} {v}" for k, v in (x.get("targetOvertimeByPerson") or {}).items()),
             f"{_d(x['testFrom'])} to {_d(x['testTo'])}", " + ".join(x["testers"]) + f", led by {x['testLead']}",
             x["appModuleCount"], x["screens"], x["ops"], x["tables"], x["hours"], x["testHours"], x["aiHours"],
             x["tasks"], x["ticketedTasks"], x["flowsClaimed"], x["flowsPartly"], ", ".join(x["criticalFlows"]),
             x["testScope"], "; ".join(x["appModules"])] for x in plan["blocks"]]
    _sheet(wb, "Blocks", ["Block", "First sprint", "Final sprint", "Ends", "Target sprint", "Target end",
                          "At normal hours", "Overtime to end on target (hours)", "By person",
                          "Block test", "Block testers", "App-modules", "Screens", "Operations", "Tables", "Build hours",
                          "Test hours", "AI engine hours", "Tasks", "Ticketed", "Flows end to end", "Flows partly testable",
                          "Revenue, safety and financial flows end to end", "Test scope", "App-modules in it"],
           rows, [6, 7, 7, 12, 7, 12, 20, 10, 50, 24, 34, 8, 8, 8, 7, 8, 8, 8, 7, 8, 8, 8, 30, 60, 120])
    write_people(wb, plan)
    rows = [[a["block"], a["name"], a["platform"], a["module"], a["package"], a["screens"], a["ops"], a["tables"],
             a["tasks"], a["ticketedTasks"], a["points"], a["hours"], a["testHours"], a["tester"],
             f"{a['sprintFrom']}-{a['sprintTo']}", a["backEndFrom"] or "", _d(a["done"]), a["lead"], a["key"]]
            for a in plan["appModules"]]
    _sheet(wb, "App-modules", ["Block", "App-module", "Platform", "Business module", "Package", "Screens", "Operations",
                               "Tables", "Tasks", "Ticketed", "Points", "Hours", "Module test hours", "Module tester",
                               "Sprint span (first-done)", "Back end from sprint", "Done", "Lead", "Feature key"],
           rows, [6, 50, 16, 28, 24, 8, 9, 7, 7, 8, 7, 7, 8, 20, 10, 9, 12, 20, 34])
    rows = [[f["id"], f["name"], f["criticality"], f["complete"] or "not in the plan",
             f["partlyFrom"], len(f["missing"]), ", ".join(f["missing"][:8])] for f in plan["flows"]]
    _sheet(wb, "Flows", ["Flow", "Name", "Criticality", "Block that completes it", "Partly testable from block",
                         "Steps' screens and operations not in the plan", "Which (first 8)"], rows,
           [7, 60, 12, 12, 12, 12, 70])
    wb.save(path)


RELEASE_CHECKLIST = [
    ("Every sprint (Friday of week 2)", "Before", "Every ticket in the sprint is Ready for QA or moved to the next sprint with a reason in the ticket.", "Scrum lead"),
    ("Every sprint", "Before", "CI green on main: unit, contract (the OpenAPI checks) and migration tests.", "Hrushikant Patkar"),
    ("Every sprint", "Before", "Database migrations run forward on a copy of staging data; the rollback script runs back cleanly.", "Back-end owner of the service"),
    ("Every sprint", "Before", "Feature flags set per tenant for anything half-built; nothing unfinished is visible to a venue.", "Lead of the package"),
    ("Every sprint", "Test", "Each app-module whose last ticket closed has its module test passed on the integration environment (docs/active/block-test-strategy.md).", "The module's tester"),
    ("Every sprint", "Deploy", "Deploy to staging; run the smoke flows of the apps touched (buy a ticket, pay, scan, POS sale, KDS bump).", "QA"),
    ("Every sprint", "Deploy", "Demo to the client from staging; record accepted and rejected items in OpenProject.", "Chinmay Parab"),
    ("Every sprint", "After", "Measured pace per person recorded; the plan regenerated with the real pace.", "Chinmay Parab"),
    ("Every block (its last three working days)", "Test", "The block test: every flow the block claims end to end, offline runs, the load run where it adds a purchase or admission path. No new feature work starts.", "The block's tester pair, led by Chinmay Parab"),
    ("Every block", "After", "The client's acceptance session on the integration environment in the following week; signed off or open items listed.", "Client"),
    ("Block A go-live", "Before", "The client has signed off every Block A wireframe batch (3 working days each, audit R252).", "Client design reviewer"),
    ("Block A go-live", "Before", "Payment sandbox credentials in place and a full sale-refund-settlement cycle passes (the client's answer in the Decisions Register).", "Client, then Tanmay Dukhande"),
    ("Block A go-live", "Before", "Offline: the POS sells and the scanner admits with the network cut, and reconcile when it returns.", "Pradnya Yeram"),
    ("Block A go-live", "Before", "Tax invoice, credit note and VAT fields checked against the client's answers (make-or-break items).", "Pranay Shinde"),
    ("Block A go-live", "Before", "Load test at the burst mix (ticket on-sale) and the venue-day mix; replicas as sized in handoff/sizing.json.", "Hrushikant Patkar"),
    ("Block A go-live", "Before", "Backups restore into a clean environment; the restore time is written down.", "Hrushikant Patkar"),
    ("Block A go-live", "Before", "Security: role grants reviewed per module, secrets in the vault, no test credentials in production.", "Hrushikant Patkar"),
    ("Block A go-live", "Deploy", "Production deploy in a quiet window; canary tenant first, then the pilot venue.", "Hrushikant Patkar"),
    ("Block A go-live", "Deploy", "Watch error rate and latency for the first hour of trading; the rollback is rehearsed beforehand.", "On-call back end"),
    ("Block A go-live", "Rollback if", "Payment failures above 2% of attempts, POS sale p95 above 2 seconds, or any failed gate admission not explained by the ticket.", "Chinmay Parab decides"),
    ("Block A go-live", "After", "Release notes to the client; tickets closed; the hypercare rota for the first two weeks.", "Chinmay Parab"),
    ("End of six months (2 Apr 2027)", "Before", "Every module's acceptance signed; open defects triaged into the phase-2 backlog.", "Chinmay Parab and the client"),
]

ARCH_DECISIONS = [
    ("One package, services by contract", "Accepted", "Each business module is one OpenAPI contract owned by one service; screens bind to operations, never to tables.",
     "Work can be split by module and checked by machine before anyone builds.", "A change crosses contract, screen and ticket; the refresh keeps them in step."),
    ("Two-week sprints, four blocks of complete app-modules", "Accepted 1 Oct", "13 sprints from 5 October to 2 April. A block is a set of app-modules (a module on one app, e.g. Ticketing · Guest Web, with the back end it needs), each complete and tested, ending on a sprint boundary: A (the first release), then B, C and D (the old B1-B3, with a step in between).",
     "Each block can be tested end to end and accepted; nothing in a block waits on a later one.", "Block A keeps its content, so it ends where its back end lands; B to D are sized to their sprints."),
    ("Every block tested end to end", "Accepted 1 Oct", "A module test per app-module (about 10% of its points, by a peer) and a block test in the last three working days of each block, when no new feature work starts (docs/active/block-test-strategy.md).",
     "A block is accepted on its flows, not on its tickets.", "About 10% more hours, and three days of each block's last sprint."),
    ("Size by formula, schedule by measured pace", "Accepted", "Points come from each screen's specification and each operation; the pace is replaced by the measured pace after the first sprints.",
     "No hand estimates; a re-plan is a rerun.", "The first forecast is only as good as Block A's planned pace."),
    ("Every AI function inside six months, baseline first", "Accepted 30 Sep", "Forecasting, fraud, recommendations, anomaly detection and both assistants ship with a working baseline on day one and learn from the tenant's own data; a trained model replaces the baseline only when it beats it and an admin approves.",
     "Customers never hear 'later, when you have data'; accuracy grows with use.", "The two AI engineers' calendar runs past 2 April (the AI review's own finding)."),
    ("Two AI engineers from 5 October", "Accepted 30 Sep", "Kalpita and the second AI engineer carry the engine work; no third engineer. Developers build the AI endpoints and screens like any other module.",
     "AI engineers spend their time on what only they can build.", "The learned producers finish in shadow after 2 April."),
    ("CMS as a flow builder", "Accepted", "Operators pick their booking flows, see required and optional steps, set their own order, then finish in a configuration panel.",
     "One engine serves every venue type the prototypes show.", "The step library has to be complete before operators can compose."),
    ("Per-tenant data and models", "Accepted", "Models train per tenant; the LLM never reads raw data (scrubbed aggregates or code it writes); retention is tenant configuration.",
     "Meets PDPL and client expectations on data isolation.", "No pooled training; cold start relies on baselines and priors."),
    ("Load follows skill; stable owners", "Accepted", "Tasks are assigned by skill rating and experience, not evenly; helpers take a share proportional to their ratings. From release r2 a pushed ticket keeps its owner; a deliberate move is a --rebalance.",
     "Hard work goes to people rated for it; boards do not reshuffle.", "Back-end owners are the Block A constraint; shown as overtime, not hidden."),
]

RISKS = [
    ("Back-end owners overloaded in Block A", "Their Block A work runs past Sprint 4 (People sheet).", "Decided 1 October: Block A is 40 working days (it was 35) and ends Sprint 4, 27 November, with the overtime the Summary sheet gives, most of it on the three back-end owners; the two new developers take back-end tasks from their first day. Its critical path is about as long as the window, so overtime cannot be bought later in the block: it has to start early. Without it, Block A ends where the Blocks sheet says at normal hours."),
    ("Pace below plan", "The plan assumes 5 tasks a developer a day through Block A, rising gradually to 2x by Sprint 11 (decided 1 October); it is re-measured after Sprint 2.", "If the measured pace is lower, rerun the plan with it; blocks B to D re-cut to their sprints. At the plan of record's 9.58 points a day, Block A needed about 996 hours of overtime to end in Sprint 4 (run of 1 October)."),
    ("Hiring the two developers slips", "Not confirmed by 23 October.", "Blocks C and D move out by the scheduler's figure; the PM decides scope or date."),
    ("Client inputs late", "Wireframe sign-off over 3 working days; sandbox credentials; stations and fares; cabana numbering; real photos.", "Those tickets wait in 'Waiting on client' and do not count against the team's pace."),
    ("Make-or-break answers", "Tax invoice fields, e-invoicing provider, VAT 201 layout, face capture consent, ID-verification provider.", "Defaults are built; a different answer is a change request."),
    ("Second AI engineer not in place on 5 October", "No start date confirmed this week.", "Kalpita starts the Block A AI alone; the baseline layer moves one sprint and the AI engine work needs more overtime."),
    ("AI engine past 2 April", "About 50 AI-engineer weeks needed against 35 available.", "Decided 1 October: every AI engine task is created; those past 2 April are unassigned until the AI developers join. Each month they do not join moves the AI engine finish by about a month."),
    ("Block test finds severity 1 or 2 defects", "A block test cannot pass in its three days.", "The defects are fixed in the next sprint's first days by the people who built the module; the block's acceptance moves, the next block does not wait."),
]


def write_md(plan, path):
    b = plan["basis"]
    a = next(x for x in plan["blocks"] if x["block"] == "A")
    L = []
    w = L.append
    w("# TICVAI complete build plan: presentation source")
    w("")
    w("> **Generated by** `tools/build-plan-deck.py` from the plan generators (the sprint plan of 1 October); rerun it after "
      "any change. The workbooks are `handoff/TICVAI - Build Plan.xlsx` and `handoff/TICVAI - Sprint Plan.xlsx`.")
    w("")
    w("Each `##` below is one slide. Numbers come from the package, not from estimates.")
    w("")
    w("## 1. The plan in one line")
    w("")
    w(f"We build TICVAI from **{_d(plan['calendar']['start'])} to {_d(plan['calendar']['planEnd'])}** in **13 two-week sprints** "
      f"and **four blocks** (A to D), each a set of complete, tested app-modules: about **{_pp(b['totalHours'])} hours** "
      f"({_pp(b['testHours'])} of them testing) across **{b['appModules']} app-modules**, with a team of {len(plan['people'])}. "
      f"At normal hours the developers finish on **{_d(b['forecastFinish'])}** and the AI engine on **{_d(b['aiFinish'])}**.")
    w("")
    w("## 2. Blocks")
    w("")
    w("| Block | Sprints | Ends | Block test | App-modules | Flows end to end | Hours (build / test) |")
    w("|---|---|---|---|---|---|---|")
    for x in plan["blocks"]:
        w(f"| {x['block']} | {x['firstSprint']}-{x['endSprint']} (target {x['targetSprint']}) | {_d(x['endsOn'])} | "
          f"{_d(x['testFrom'])} to {_d(x['testTo'])} | {x['appModuleCount']} | {x['flowsClaimed']} (+{x['flowsPartly']} partly) | "
          f"{_pp(x['hours'])} / {_pp(x['testHours'])} |")
    w("")
    w("## 3. Sprints")
    w("")
    w("| Sprint | Dates | Blocks | Capacity (h) | Planned (h) | Block test |")
    w("|---|---|---|---|---|---|")
    for s in plan["sprints"]:
        bt = "; ".join(f"{x['block']}: {_d(x['from'])}-{_d(x['to'])}" for x in s["blockTest"])
        w(f"| {s['n']}{' (buffer)' if s.get('buffer') else ''} | {_d(s['start'])} – {_d(s['end'])} | {', '.join(s['blocks'])} | {_pp(s['capacity'])} | {_pp(s['planned'])} | {bt} |")
    w("")
    w("Holidays counted: " + ("; ".join(f"{_d(d)} {n}" for d, n in plan["calendar"]["holidays"].items())
                              or "none: every weekday is a working day, and leave is handled when it comes up") + ".")
    w("")
    w("## 4. Packages")
    w("")
    w("| Package | App-modules | Requirements | Screens | Operations | Hours | Completion | Lead |")
    w("|---|---|---|---|---|---|---|---|")
    for p in plan["packages"]:
        w(f"| **{p['package']}** | {p['appModules']} | {_pp(p['requirements'])} | {_pp(p['screens'])} | {_pp(p['ops'])} | {_pp(p['hours'])} | {_d(p['end'])} | {p['lead']} |")
    w("")
    w("## 5. Build phases, end to end")
    w("")
    for side in ("back end", "front end"):
        w(f"**{side.capitalize()}**")
        w("")
        w("| Phase | Starts | Ends | Hours (A / B / C / D) | Main modules | People |")
        w("|---|---|---|---|---|---|")
        for r in plan["phases"]:
            if r["side"] != side:
                continue
            w(f"| {r['phase']} {r['name']} | {_d(r['start'])} | {_d(r['end'])} | "
              + " / ".join(_pp(r["hoursByBlock"][x]) for x in sp.BLOCKS)
              + f" | {', '.join(r['modules'][:5])} | {', '.join(r['people'][:4])} |")
        w("")
    w("## 6. The team")
    w("")
    w("| Name | Role | Block A points | Block A work ends | Last day of planned work | Hours past 2 April |")
    w("|---|---|---|---|---|---|")
    for pe in plan["people"]:
        w(f"| {pe['name']} | {pe['role']} | {pe['blockAPoints'] or '–'} | {_d(pe['blockAEnd']) or '–'} | {_d(pe['lastDay'])} | {pe['hoursAfterPlanEnd']} |")
    w("")
    w(b["notCounted"] + " Loads are uneven on purpose: they follow skill and experience.")
    w("")
    w("## 7. Testing")
    w("")
    w(f"- **Module tests:** one per app-module, about 10% of its points, by a peer who did not build most of it: {_pp(b['moduleTestHours'])} hours.")
    w(f"- **Block tests:** the last three working days of each block, a back-end and front-end pair led by Chinmay Parab: {_pp(b['blockTestHours'])} hours. No new feature work starts in those days.")
    w("- **Flows:** each block is accepted on the flows it completes (the Flows sheet).")
    w("")
    w("## 8. Finishing by 2 April: overtime")
    w("")
    w(f"About **{_pp(b['overtimeHoursDevelopers'])} developer hours** past 2 April at normal hours. Decided 1 October: Block D keeps its scope; "
      f"{d_text(b)} Block A ends Sprint {a['targetSprint']} with {_pp(_opt_a(b, a)['overtimeHours'])} hours of planned overtime ("
      + ", ".join(f"{n} {h}" for n, h in _opt_a(b, a)['byPerson'].items()) + f"). {buffer_text(b)} Ticketing depth: {TICKETING} "
      f"The AI engine's {_pp(b.get('aiUnassignedHours', 0))} hours past 2 April are not overtime: those tasks are unassigned for the AI developers joining. "
      f"The plan of {PREVIOUS['date']} needed about {_pp(PREVIOUS['overtimeHours'])} hours, without the module and block tests.")
    w("")
    w("## 9. Architecture decisions the plan rests on")
    w("")
    for x in ARCH_DECISIONS:
        w(f"- **{x[0]}** ({x[1]}). {x[2]}")
    w("")
    w("## 10. Risks")
    w("")
    for r_ in RISKS:
        w(f"- **{r_[0]}.** Signal: {r_[1]} Action: {r_[2]}")
    w("")
    io.open(path, "w", encoding="utf-8").write("\n".join(L) + "\n")


if __name__ == "__main__":
    plan_, task_rows_ = main()
    plan = json.loads(json.dumps(plan_, default=jsonable, ensure_ascii=False))

    def o(rel):
        return os.path.join(OUT, rel) if OUT != "." else rel
    io.open(o("handoff/build-plan.json"), "w", encoding="utf-8").write(json.dumps(plan, indent=1, ensure_ascii=False))
    write_build_xlsx(plan, o("handoff/TICVAI - Build Plan.xlsx"))
    write_sprint_xlsx(plan, task_rows_, o("handoff/TICVAI - Sprint Plan.xlsx"))
    write_md(plan, o("docs/active/build-plan-presentation.md"))
    b = plan["basis"]
    print(f"  pace: {b.get('pace')}; total {b['totalHours']} h "
          f"(build {b['buildHours']}, tests {b['testHours']}, AI engine {b['aiEngineHours']}); developers finish "
          f"{b['forecastFinish']}, AI engine {b['aiFinish']}; overtime to 2 Apr {b['overtimeHours']} h")
    for x in plan["blocks"]:
        print(f"  Block {x['block']}: Sprints {x['firstSprint']}-{x['endSprint']} (target {x['targetSprint']}), ends "
              f"{x['endsOn']}, {x['appModuleCount']} app-modules, {x['flowsClaimed']} flows end to end "
              f"(+{x['flowsPartly']} partly)")
    print("  -> handoff/build-plan.json, handoff/TICVAI - Build Plan.xlsx, handoff/TICVAI - Sprint Plan.xlsx, "
          "docs/active/build-plan-presentation.md")
